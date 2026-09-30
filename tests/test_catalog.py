"""Generic tests for any template pack's layout catalog: uniqueness, target validity, buildability."""

from __future__ import annotations

import copy
from typing import Any

import pytest
from pptx import Presentation

from deckwright import pack
from deckwright.engine import DeckBuilder
from deckwright.pack import EXAMPLE_IMAGE, Template

PACK_IDS = ["sample"]


@pytest.fixture(params=PACK_IDS, ids=PACK_IDS)
def any_pack(request) -> Template:
    return pack.load(request.param)


def _replace_img(value: Any, replacement: str) -> Any:
    """Recursively replace the placeholder image url with a local test image path."""
    if isinstance(value, str):
        return replacement if value == EXAMPLE_IMAGE else value
    if isinstance(value, list):
        return [_replace_img(v, replacement) for v in value]
    if isinstance(value, dict):
        return {k: _replace_img(v, replacement) for k, v in value.items()}
    return value


def test_layout_ids_are_unique(any_pack):
    ids = [layout.id for layout in any_pack.layouts]
    assert len(ids) == len(set(ids))
    assert len(any_pack.by_id) == len(any_pack.layouts)


def _all_shape_ids(shapes) -> set[int]:
    ids: set[int] = set()
    for shape in shapes:
        ids.add(shape.shape_id)
        if shape.shape_type == 6:  # MSO_SHAPE_TYPE.GROUP
            ids |= _all_shape_ids(shape.shapes)
    return ids


def _all_placeholder_idxs(slide_layout) -> set[int]:
    return {ph.placeholder_format.idx for ph in slide_layout.placeholders}


def _layout_targets(layout) -> list[int]:
    targets = [f.target for f in layout.fields]
    targets += list(layout.strip)
    if layout.items:
        for lst in layout.items.targets.values():
            targets += lst
        for lst in layout.items.extra.values():
            targets += lst
    return targets


def test_targets_exist_on_their_source(any_pack):
    prs = Presentation(str(any_pack.pptx))
    for layout in any_pack.layouts:
        if layout.source == "slide":
            slide = prs.slides[layout.ref - 1]
            shape_ids = _all_shape_ids(slide.shapes)
            for target in _layout_targets(layout):
                assert target in shape_ids, f"{layout.id}: shape id {target} not found on slide {layout.ref}"
        else:
            slide_layout = prs.slide_layouts[layout.ref]
            idxs = _all_placeholder_idxs(slide_layout)
            for target in _layout_targets(layout):
                assert target in idxs, f"{layout.id}: placeholder idx {target} not found in layout {layout.ref}"


def test_every_layout_builds_from_its_own_example(any_pack, img):
    for layout in any_pack.layouts:
        content = _replace_img(copy.deepcopy(layout.example), img)
        builder = DeckBuilder(any_pack)
        builder.add(any_pack.by_id[layout.id], content)
        result = builder.build()
        assert result.data.startswith(b"PK"), layout.id


def test_sample_pack_examples_build_with_zero_warnings(img):
    t = pack.load("sample")
    for layout in t.layouts:
        content = _replace_img(copy.deepcopy(layout.example), img)
        builder = DeckBuilder(t)
        builder.add(t.by_id[layout.id], content)
        result = builder.build()
        assert not result.warnings, (layout.id, result.warnings)
