"""Tests for deckwright.engine.DeckBuilder: slide assembly, text fill, warnings."""

from __future__ import annotations

import io
from pathlib import Path

from pptx import Presentation
from pptx.oxml.ns import qn

from deckwright.deck import build_deck
from deckwright.engine import DeckBuilder, _shape_by_id
from deckwright.models import DeckSpec


def _prs_from_bytes(data: bytes) -> Presentation:
    return Presentation(io.BytesIO(data))


def _all_shapes(shapes):
    for shape in shapes:
        yield shape
        if shape.shape_type == 6:  # GROUP
            yield from _all_shapes(shape.shapes)


def _all_text(prs) -> str:
    chunks = []
    for slide in prs.slides:
        for shape in _all_shapes(slide.shapes):
            if shape.has_text_frame:
                chunks.append(shape.text_frame.text)
    return "\n".join(chunks)


def _shape_el_by_id(slide, shape_id: int):
    return _shape_by_id(slide.shapes._spTree, shape_id)


def test_deck_of_three_slides_has_exactly_three_slides(template):
    builder = DeckBuilder(template)
    builder.add(template.by_id["title"], {"title": "Hello", "subtitle": "2026", "presenter": "Me"})
    builder.add(template.by_id["statement"], {"title": "A **point**"})
    builder.add(template.by_id["closing"], {})
    result = builder.build()
    prs = _prs_from_bytes(result.data)
    assert len(prs.slides) == 3


def test_highlight_markup_produces_pink_run(template):
    builder = DeckBuilder(template)
    layout = template.by_id["title"]  # title field target 3, highlight=True
    builder.add(layout, {"title": "Hello **word** there", "subtitle": "d", "presenter": "p"})
    result = builder.build()
    prs = _prs_from_bytes(result.data)
    slide = prs.slides[0]
    el = _shape_el_by_id(slide, 3)
    assert el is not None
    found = False
    for r in el.iter(qn("a:r")):
        t = r.find(qn("a:t"))
        if t is not None and t.text == "word":
            rpr = r.find(qn("a:rPr"))
            assert rpr is not None
            fill = rpr.find(qn("a:solidFill"))
            assert fill is not None
            scheme = fill.find(qn("a:schemeClr"))
            assert scheme is not None
            assert scheme.get("val") == "accent1"
            found = True
    assert found, "expected a run with text 'word'"


def test_speaker_notes_land_on_notes_slide(template):
    builder = DeckBuilder(template)
    builder.add(template.by_id["statement"], {"title": "Point"}, notes="Remember to smile.")
    result = builder.build()
    prs = _prs_from_bytes(result.data)
    slide = prs.slides[0]
    assert slide.has_notes_slide
    assert "Remember to smile." in slide.notes_slide.notes_text_frame.text


def test_agenda_with_two_items_removes_unused_shapes_and_numbers_them(template):
    builder = DeckBuilder(template)
    layout = template.by_id["agenda"]
    builder.add(layout, {"items": [{"label": "First"}, {"label": "Second"}]})
    result = builder.build()
    prs = _prs_from_bytes(result.data)
    slide = prs.slides[0]

    number_ids = layout.items.targets["number"]
    label_ids = layout.items.targets["label"]

    # positions 2 and 3 (unused) should be gone
    assert _shape_el_by_id(slide, number_ids[2]) is None
    assert _shape_el_by_id(slide, number_ids[3]) is None
    assert _shape_el_by_id(slide, label_ids[2]) is None
    assert _shape_el_by_id(slide, label_ids[3]) is None

    number0 = _shape_el_by_id(slide, number_ids[0])
    number1 = _shape_el_by_id(slide, number_ids[1])
    assert number0 is not None and number1 is not None

    def texts(el):
        return "".join(t.text or "" for t in el.iter(qn("a:t")))

    assert texts(number0) == "01"
    assert texts(number1) == "02"


def test_missing_required_image_warns_about_pink_placeholder(template):
    builder = DeckBuilder(template)
    layout = template.by_id["text-image"]  # image field is required
    builder.add(layout, {"title": "No image here", "body": "b"})
    result = builder.build()
    assert any("pink placeholder" in w for w in result.warnings), result.warnings


def test_bar_chart_width_grows_with_percentage(template):
    builder = DeckBuilder(template)
    layout = template.by_id["bar-chart"]
    builder.add(layout, {"title": "T", "items": [{"value": "20%"}, {"value": "80%"}]})
    result = builder.build()
    prs = _prs_from_bytes(result.data)
    slide = prs.slides[0]
    bar_ids = layout.items.extra["bar"]
    bar0 = _shape_el_by_id(slide, bar_ids[0])
    bar1 = _shape_el_by_id(slide, bar_ids[1])
    assert bar0 is not None and bar1 is not None

    def width(el):
        ext = el.find(qn("p:spPr")).find(qn("a:xfrm")).find(qn("a:ext"))
        return int(ext.get("cx"))

    assert width(bar0) < width(bar1)


def test_very_long_text_in_small_box_warns(template):
    builder = DeckBuilder(template)
    layout = template.by_id["statement"]  # very large font, one sentence
    long_text = " ".join(["supercalifragilisticexpialidocious"] * 60)
    builder.add(layout, {"title": long_text})
    result = builder.build()
    assert any(("shrunk" in w or "too long" in w) for w in result.warnings), result.warnings


def test_code_handler_fills_line_numbers_matching_line_count(template):
    builder = DeckBuilder(template)
    layout = template.by_id["code"]
    code = "line one\nline two\nline three"
    builder.add(layout, {"title": "T", "code": code})
    result = builder.build()
    prs = _prs_from_bytes(result.data)
    slide = prs.slides[0]
    numbers_el = _shape_el_by_id(slide, layout.handler.line_numbers)
    assert numbers_el is not None
    labels = [t.text for t in numbers_el.iter(qn("a:t"))]
    assert labels == ["01", "02", "03"]


def test_footer_logo_on_master_layout_slide_present_and_absent(template):
    layout = template.by_id["title-content"]

    with_footer = build_deck(DeckSpec(slides=[{"layout": layout.id, "title": "T", "body": "b"}]), template)
    slide = _prs_from_bytes(with_footer.data).slides[0]
    images = [r for r in slide.part.rels.values() if r.reltype.endswith("/image")]
    assert len(images) == 1

    without_footer = build_deck(
        DeckSpec(slides=[{"layout": layout.id, "title": "T", "body": "b"}], footer=False), template
    )
    slide2 = _prs_from_bytes(without_footer.data).slides[0]
    images2 = [r for r in slide2.part.rels.values() if r.reltype.endswith("/image")]
    assert len(images2) == 0


def test_raw_mode_layout_index_and_placeholders(template):
    # sample layout 0 ("Title and Content") has placeholder idx 0 (title) and 1 (body)
    spec = DeckSpec.model_validate({
        "slides": [{"layout_index": 0, "placeholders": {"0": "Raw"}}]
    })
    result = build_deck(spec, template)
    prs = _prs_from_bytes(result.data)
    assert len(prs.slides) == 1
    assert "Raw" in _all_text(prs)


def test_sample_deck_example_has_no_lorem_or_placeholder_text():
    repo_root = Path(__file__).resolve().parents[1]
    spec_path = repo_root / "examples" / "sample-deck.json"
    spec = DeckSpec.model_validate_json(spec_path.read_text())
    result = build_deck(spec)
    prs = _prs_from_bytes(result.data)
    text = _all_text(prs)
    assert "Lorem" not in text
    assert not result.warnings, result.warnings
