"""Regression tests for the code-review findings."""

from __future__ import annotations

import io
import json
import shutil
import zipfile

import pytest
from pptx import Presentation

from deckwright.deck import build_deck
from deckwright.engine import TemplateError, _percent
from deckwright.images import ImageSourceError, load_image
from deckwright.models import DeckSpec


def _media(data: bytes) -> list[str]:
    return [n for n in zipfile.ZipFile(io.BytesIO(data)).namelist() if n.startswith("ppt/media/")]


def test_replaced_template_photos_are_not_shipped(img):
    one = build_deck(DeckSpec(slides=[{"layout": "text-image", "title": "T", "image": img}]))
    prs = Presentation(io.BytesIO(one.data))
    slide = prs.slides[0]
    images = [r for r in slide.part.rels.values() if r.reltype.endswith("/image")]
    # the user image plus the footer logo, nothing else
    assert len(images) == 2
    referenced = {el.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed")
                  for el in slide._element.iter()}
    assert all(r.rId in referenced for r in images)


def test_removed_item_image_not_shipped(img):
    # images-2 has two picture slots; giving only one item replaces the second
    # template photo with a placeholder, so the template's own photo isn't shipped.
    data = build_deck(DeckSpec(slides=[{"layout": "images-2", "items": [{"image": img}]}])).data
    slide = Presentation(io.BytesIO(data)).slides[0]
    rids = {r.rId for r in slide.part.rels.values() if r.reltype.endswith("/image")}
    ns = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"
    used = {el.get(ns) for el in slide._element.iter() if el.get(ns)}
    assert rids == used
    assert len(rids) == 3  # provided photo, placeholder for the missing item, footer logo


def test_small_deck_is_much_smaller_than_template(img):
    data = build_deck(DeckSpec(slides=[{"layout": "statement", "title": "Hi"}])).data
    assert len(_media(data)) < 20


@pytest.mark.parametrize(
    ("value", "expected"),
    [("84%", 84), ("1,200", 1200), ("1.200,5", 1200.5), ("12,5%", 12.5), ("-3", -3), ("n/a", 0), ("48.", 48)],
)
def test_percent_parsing(value, expected):
    assert _percent(value) == expected


def test_explicit_null_uses_default():
    data = build_deck(DeckSpec(slides=[{"layout": "agenda", "title": None, "items": [{"label": "A"}]}])).data
    texts = [s.text_frame.text for s in Presentation(io.BytesIO(data)).slides[0].shapes if s.has_text_frame]
    assert "Agenda" in texts


def test_raw_mode_layout_index_and_placeholders():
    # sample layout 0 ("Title and Content") has placeholder idx 0 (title) and 1 (body)
    spec = DeckSpec.model_validate({"slides": [{"layout_index": 0, "placeholders": {"0": "Raw"}}]})
    data = build_deck(spec).data
    prs = Presentation(io.BytesIO(data))
    assert len(prs.slides) == 1
    texts = [s.text_frame.text for s in prs.slides[0].shapes if s.has_text_frame]
    assert "Raw" in texts


def test_raw_mode_rejects_non_numeric_key():
    with pytest.raises(TemplateError):
        build_deck(DeckSpec(slides=[{"layout_index": 0, "placeholders": {"title": "x"}}]))


def test_api_raw_mode_bad_key_is_422(output_dir):
    from fastapi.testclient import TestClient

    import deckwright.api as api

    r = TestClient(api.app).post("/v1/presentations", json={"slides": [{"layout_index": 0, "placeholders": {"t": "x"}}]})
    assert r.status_code == 422


@pytest.mark.parametrize("url", ["http://127.0.0.1/x.png", "http://localhost/x.png", "http://169.254.169.254/latest"])
def test_private_urls_refused_for_untrusted_callers(url):
    with pytest.raises(ImageSourceError):
        load_image(url, allow_local=False)


def test_api_refuses_private_url(output_dir):
    from fastapi.testclient import TestClient

    import deckwright.api as api

    spec = {"slides": [{"layout": "text-image", "title": "T", "image": "http://127.0.0.1:9/x.png"}]}
    warnings = TestClient(api.app).post("/v1/presentations", json=spec).json()["warnings"]
    assert any("non-public" in w for w in warnings)


def test_file_endpoint_returns_warning_json(output_dir):
    from fastapi.testclient import TestClient

    import deckwright.api as api

    r = TestClient(api.app).post("/v1/presentations.pptx", json={"slides": [{"layout": "text-image", "title": "T"}]})
    assert r.status_code == 200
    warnings = json.loads(r.headers["x-deckwright-warnings"])
    assert int(r.headers["x-deckwright-warning-count"]) == len(warnings) >= 1


@pytest.mark.skipif(shutil.which("pdftoppm") is None, reason="needs poppler and LibreOffice")
def test_preview_is_cached(output_dir):
    import deckwright.service as service

    out = service.create(DeckSpec(slides=[{"layout": "statement", "title": "A"}, {"layout": "closing"}]))
    first = service.preview(out["id"], 1, 1, dpi=20)
    mtime = first[0].stat().st_mtime
    second = service.preview(out["id"], 2, 2, dpi=20)
    assert len(first) == len(second) == 1
    assert service.preview(out["id"], 1, 1, dpi=20)[0].stat().st_mtime == mtime
