"""Placeholder images and Excalidraw diagrams."""

from __future__ import annotations

import io
import json
from pathlib import Path

import pytest
from PIL import Image
from pptx import Presentation

from deckwright import diagram
from deckwright.brand import DIAGRAM_GUIDE
from deckwright.deck import build_deck
from deckwright.models import DeckSpec

SCENE = DIAGRAM_GUIDE["example"]["excalidraw"]
EXAMPLES = Path(__file__).resolve().parents[1] / "examples"


@pytest.fixture
def style(template):
    return diagram.Style.from_template(template)


def _picture_sizes(data: bytes, slide: int = 0) -> list[tuple[int, int]]:
    prs = Presentation(io.BytesIO(data))
    out = []
    for shape in prs.slides[slide].shapes:
        if shape._element.tag.endswith("}pic") and shape.width > 2 * 914400:  # not the footer logo
            out.append(Image.open(io.BytesIO(shape.image.blob)).size)
    return out


def test_placeholder_string_and_dict_make_todos_and_notes(template):
    spec = DeckSpec(slides=[
        {"layout": "text-image", "title": "W", "body": "b", "image": "placeholder: Screenshot of the canvas"},
        {"layout": "images-2", "items": [{"image": {"placeholder": "Left photo"}}, {"image": "placeholder: Right"}]},
    ])
    r = build_deck(spec, template)
    assert [t["label"] for t in r.todos] == ["Screenshot of the canvas", "Left photo", "Right"]
    notes = Presentation(io.BytesIO(r.data)).slides[0].notes_slide.notes_text_frame.text
    assert "TODO" in notes and "Screenshot of the canvas" in notes
    assert not r.warnings


def test_placeholder_matches_frame_aspect(template):
    r = build_deck(DeckSpec(slides=[{"layout": "text-image", "title": "W", "body": "b",
                                     "image": "placeholder: X"}]), template)
    (w, h), = _picture_sizes(r.data)
    assert abs(w / h - 5.9 / 5.6) < 0.05  # text-image slot is about 5.9 x 5.6 in


def test_placeholder_keeps_user_notes(template):
    r = build_deck(DeckSpec(slides=[{"layout": "text-image", "title": "W", "body": "b",
                                     "image": "placeholder: X", "notes": "Say hi"}]), template)
    notes = Presentation(io.BytesIO(r.data)).slides[0].notes_slide.notes_text_frame.text
    assert notes.startswith("Say hi") and "TODO" in notes


def test_diagram_renders_and_saves_editable_file(template):
    r = build_deck(DeckSpec(slides=[{"layout": "text-image", "title": "W", "body": "b",
                                     "image": {"excalidraw": SCENE}}]), template)
    assert not r.warnings
    assert list(r.assets) == ["slide01-image.excalidraw"]
    saved = json.loads(r.assets["slide01-image.excalidraw"])
    assert saved["type"] == "excalidraw" and saved["version"] == 2
    texts = [e for e in saved["elements"] if e["type"] == "text"]
    assert {t["text"] for t in texts} == {"Webhook", "AI agent", "Route ticket"}
    assert all(t["containerId"] for t in texts)
    assert _picture_sizes(r.data)


def test_diagram_from_excalidraw_file(tmp_path, template):
    f = tmp_path / "flow.excalidraw"
    f.write_text(json.dumps(SCENE))
    r = build_deck(DeckSpec(slides=[{"layout": "text-image", "title": "W", "body": "b", "image": str(f)}]), template)
    assert not r.warnings and r.assets


def test_diagram_file_refused_for_untrusted_callers(tmp_path, template):
    f = tmp_path / "flow.excalidraw"
    f.write_text(json.dumps(SCENE))
    r = build_deck(DeckSpec(slides=[{"layout": "text-image", "title": "W", "body": "b", "image": str(f)}]), template,
                   allow_local_files=False)
    assert any("disabled" in w for w in r.warnings)


def test_bad_diagram_falls_back_with_warning(template):
    r = build_deck(DeckSpec(slides=[{"layout": "text-image", "title": "W", "body": "b",
                                     "image": {"excalidraw": {"elements": "nope"}}}]), template)
    assert any("could not be rendered" in w for w in r.warnings)


def test_brand_colour_mapping(style):
    assert diagram.brand_color("#1e1e1e", style) == style.black
    assert diagram.brand_color("#e03131", style) in (style.accent, style.accent_dark)
    assert diagram.brand_color("#a5d8ff", style) in style.accent_tints
    assert diagram.brand_color("transparent", style) is None


def test_render_is_transparent_png_of_frame_size(style):
    png = diagram.render(SCENE, 10, 5, style)
    im = Image.open(io.BytesIO(png))
    assert im.size == (1000, 500) and im.mode == "RGBA"
    assert im.getpixel((0, 0))[3] == 0


# ---- code-review regressions


def test_excalidraw_dict_with_path_string_does_not_read_files(tmp_path, template):
    f = tmp_path / "secret.excalidraw"
    f.write_text(json.dumps(SCENE))
    r = build_deck(DeckSpec(slides=[{"layout": "text-image", "title": "W", "body": "b",
                                     "image": {"excalidraw": str(f)}}]), template, allow_local_files=False)
    assert not r.assets
    assert any("could not be rendered" in w for w in r.warnings)


def test_scene_limits():
    with pytest.raises(ValueError):
        diagram.normalise({"elements": [{"type": "rectangle"}] * (diagram.MAX_ELEMENTS + 1)})
    with pytest.raises(ValueError):
        diagram.normalise({"elements": [{"type": "freedraw", "points": [[i, i] for i in range(3000)]}]})


def test_negative_size_boxes_are_normalised(style):
    el = diagram.normalise({"elements": [{"type": "rectangle", "x": 100, "y": 50, "width": -80, "height": -40}]})
    r = el["elements"][0]
    assert (r["x"], r["y"], r["width"], r["height"]) == (20, 10, 80, 40)
    diagram.render({"elements": [{"type": "ellipse", "x": 0, "y": 0, "width": -50, "height": 30}]}, 4, 3, style)


def test_unsized_text_is_measured():
    t = diagram.normalise({"elements": [{"type": "text", "x": 0, "y": 0, "text": "a long label here", "fontSize": 20}]})
    assert t["elements"][0]["width"] > 150


def test_arrow_label_sits_at_the_middle():
    scene = diagram.normalise({"elements": [
        {"type": "arrow", "id": "a", "x": 0, "y": 0, "points": [[0, 0], [300, 0]], "label": {"text": "yes"}}]})
    label = next(e for e in scene["elements"] if e["type"] == "text")
    assert label["containerId"] is None
    assert abs(label["x"] + label["width"] / 2 - 150) < 1


def test_unknown_dict_source_warns_clearly(template):
    r = build_deck(DeckSpec(slides=[{"layout": "text-image", "title": "W", "body": "b",
                                     "image": {"url": "https://x"}}]), template)
    assert any("image source has keys ['url']" in w for w in r.warnings)


def test_translucent_elements_still_render(style):
    png = diagram.render({"elements": [{"type": "rectangle", "x": 0, "y": 0, "width": 100, "height": 100,
                                        "backgroundColor": "#2F6FEB", "opacity": 50}]}, 2, 2, style)
    im = Image.open(io.BytesIO(png))
    r, g, b, a = im.getpixel((100, 100))
    assert 100 < a < 160


def test_diagrams_and_placeholders_example_builds_clean(template):
    """The bundled example mixes string/dict placeholders and Excalidraw diagrams; it must build warning-free."""
    spec = DeckSpec.model_validate_json((EXAMPLES / "diagrams-and-placeholders.json").read_text())
    r = build_deck(spec, template)
    assert r.data.startswith(b"PK")
    assert not r.warnings, r.warnings
    excalidraw_assets = [name for name in r.assets if name.endswith(".excalidraw")]
    assert len(excalidraw_assets) == 2
    for name in excalidraw_assets:
        saved = json.loads(r.assets[name])
        assert saved["type"] == "excalidraw" and saved["elements"]
