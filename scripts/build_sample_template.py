"""Build the neutral sample template pack that ships with Deckwright.

Writes src/deckwright/templates/sample/ (template.pptx, pack.json, fonts/).
The template is 16:9, blue accent, Manrope, with one hidden guide slide,
designed sample slides for every common kind and three master-layout designs.
The pack is declared here from the shapes this script creates. The generator
must later reproduce it from template.pptx alone.

Run: uv run python scripts/build_sample_template.py
Then: uv run python scripts/build_thumbnails.py --template sample
"""

from __future__ import annotations

import io
import re
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_THEME_COLOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

from deckwright.pack import (
    TEMPLATE_FILE,
    BarsHandler,
    Brand,
    CodeHandler,
    FieldSpec,
    Footer,
    Highlight,
    ItemsSpec,
    LayoutSpec,
    Pack,
    Palette,
    sha256,
    write_pack,
)

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "src" / "deckwright" / "templates" / "sample"
FONT_FILES = ["Manrope-Regular.ttf", "Manrope-Bold.ttf", "ManropeMedium-Regular.ttf", "ManropeMedium-Bold.ttf",
              "GeistMonoSemiBold-Regular.ttf", "GeistMonoSemiBold-Bold.ttf"]

W, H = 13.333, 7.5
M = 0.8  # side margin, inches
FONT, MEDIUM, MONO = "Manrope", "Manrope Medium", "Geist Mono SemiBold"
PALETTE = Palette(text="1F2328", background="FFFFFF", accent="2F6FEB",
                  accent_tints=("EAF1FE", "D3E2FD", "A9C6FB"), accent_dark="1A4DB3", accent_darkest="0B2A66")
GREY = "F2F3F5"
LOREM = "One or two short sentences that support the point on this slide."


def example_images(folder: Path, package_image: Path) -> None:
    """Neutral gradient images for examples and the sample deck (no photos, no licence questions)."""
    folder.mkdir(parents=True, exist_ok=True)
    top, bottom = (0x1F, 0x23, 0x28), (0x2F, 0x6F, 0xEB)
    for name, (w, h) in {"landscape": (1600, 900), "portrait": (900, 1200), "square": (800, 800),
                         "wide": (2400, 800)}.items():
        img = Image.new("RGB", (w, h))
        for y in range(h):
            t = y / (h - 1)
            img.paste(tuple(int(a + (b - a) * t) for a, b in zip(top, bottom)), (0, y, w, y + 1))
        r = min(w, h) // 4
        from PIL import ImageDraw

        ImageDraw.Draw(img).ellipse((w // 2 - r, h // 2 - r, w // 2 + r, h // 2 + r), outline=(0xA9, 0xC6, 0xFB),
                                    width=max(4, r // 25))
        img.save(folder / f"{name}.jpg", quality=85)
        if name == "landscape":
            img.save(package_image, quality=82)


def rgb(hex_: str) -> RGBColor:
    return RGBColor.from_string(hex_)


# --------------------------------------------------------------------------- theme and masters


def _theme(prs: Presentation) -> None:
    """Accent colours and fonts in the theme, so inherited text is on-brand."""
    part = prs.slide_master.part.part_related_by(RT.THEME)
    xml = part.blob.decode()
    colors = {"dk1": PALETTE.text, "lt1": "FFFFFF", "dk2": PALETTE.accent_darkest, "lt2": GREY,
              "accent1": PALETTE.accent, "accent2": PALETTE.accent_tints[0], "accent3": PALETTE.accent_tints[1],
              "accent4": PALETTE.accent_tints[2], "accent5": PALETTE.accent_dark, "accent6": PALETTE.accent_darkest}
    for slot, val in colors.items():
        xml = re.sub(rf"<a:{slot}>.*?</a:{slot}>", f'<a:{slot}><a:srgbClr val="{val}"/></a:{slot}>', xml, flags=re.S)
    xml = re.sub(r'(<a:(?:major|minor)Font>\s*<a:latin typeface=")[^"]*"', rf'\g<1>{FONT}"', xml)
    part._blob = xml.encode()


def _widen_masters(prs: Presentation, factor: float) -> None:
    """The default template is 4:3. Stretch master and layout boxes horizontally to 16:9."""
    parts = [prs.slide_master, *prs.slide_layouts]
    for p in parts:
        for tag in ("a:off", "a:ext"):
            for el in p._element.iter(qn(tag)):
                key = "x" if tag == "a:off" else "cx"
                if el.get(key) is not None:
                    el.set(key, str(int(int(el.get(key)) * factor)))


# --------------------------------------------------------------------------- shape helpers


def text(slide, x, y, w, h, value, size, *, font=FONT, color=None, bold=False, align=None, anchor=None,
         runs=None):
    """Text box. ``runs`` is a list of (text, is_accent) for mixed colour."""
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.word_wrap = True
    if anchor:
        tf.vertical_anchor = anchor
    paras = value if isinstance(value, list) else [value]
    for i, ptext in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        if align:
            p.alignment = align
        for rtext, accent in (runs if runs and i == 0 else [(ptext, False)]):
            r = p.add_run()
            r.text = rtext
            r.font.size = Pt(size)
            r.font.name = font
            r.font.bold = bold
            if accent:
                r.font.color.theme_color = MSO_THEME_COLOR.ACCENT_1
            elif color:
                r.font.color.rgb = rgb(color)
    return shape


def box(slide, x, y, w, h, fill, shape=MSO_SHAPE.RECTANGLE):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = rgb(fill)
    s.line.fill.background()
    style = s._element.find(qn("p:style"))  # drop the theme line, fill and shadow references
    if style is not None:
        s._element.remove(style)
    return s


def picture(slide, x, y, w, h):
    shade = Image.new("RGB", (int(w * 40), int(h * 40)), (0xD0, 0xD4, 0xDA))
    buf = io.BytesIO()
    shade.save(buf, "PNG")
    buf.seek(0)
    return slide.shapes.add_picture(buf, Inches(x), Inches(y), Inches(w), Inches(h))


def _logo_png() -> io.BytesIO:
    img = Image.new("RGBA", (120, 120), (0, 0, 0, 0))
    for i in range(3):
        for j in range(3):
            if (i + j) % 2 == 0:
                img.paste((0x2F, 0x6F, 0xEB, 255), (i * 40 + 4, j * 40 + 4, i * 40 + 36, j * 40 + 36))
    buf = io.BytesIO()
    img.save(buf, "PNG")
    buf.seek(0)
    return buf


LOGO = _logo_png().getvalue()


def footer(slide):
    """Small logo at the bottom left: the same picture on every design slide."""
    return slide.shapes.add_picture(io.BytesIO(LOGO), Inches(M), Inches(H - 0.65), Inches(0.35), Inches(0.35))


def background(slide, hex_):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = rgb(hex_)


def T(name, shape, **kw):
    return FieldSpec(name=name, target=shape.shape_id, **kw)


def I(name, shape, **kw):  # noqa: E743
    return FieldSpec(name=name, target=shape.shape_id, type="image", **kw)


def ids(shapes):
    return [s.shape_id for s in shapes]


# --------------------------------------------------------------------------- designs


def build() -> tuple[Presentation, list[LayoutSpec], Footer]:
    prs = Presentation()
    factor = Inches(W) / prs.slide_width
    prs.slide_width, prs.slide_height = Inches(W), Inches(H)
    _widen_masters(prs, factor)
    _theme(prs)
    keep = {n: prs.slide_layouts[i] for n, i in (("content", 1), ("section", 2), ("blank", 6), ("picture", 8))}
    for lay in [lay for lay in prs.slide_layouts if lay not in keep.values()]:
        prs.slide_layouts.remove(lay)  # only the layouts the pack declares stay in the template
    blank = keep["blank"]

    def ref(name):
        return prs.slide_layouts.index(keep[name])
    layouts: list[LayoutSpec] = []
    footer_ref: Footer | None = None

    def new(dark=False):
        nonlocal footer_ref
        s = prs.slides.add_slide(blank)
        if dark:
            background(s, PALETTE.text)
        logo = footer(s)
        if footer_ref is None:  # the first design slide holds the footer that master-layout slides copy
            footer_ref = Footer(slide=len(prs.slides), shape=logo.shape_id)
        return s, logo

    def add(**kw):
        layouts.append(LayoutSpec(source="slide", ref=len(prs.slides), **kw))

    # 1: hidden guide slide
    g = prs.slides.add_slide(blank)
    text(g, M, 0.6, 11, 0.8, "Brand guide", 32, font=MEDIUM)
    text(g, M, 1.6, 11, 4.5, [
        "Use one accent colour. Mark a few key words with **double asterisks**.",
        "Keep card titles under 10 words.",
        "Vary layouts. Do not put every section on the same text slide.",
    ], 18)
    g._element.set("show", "0")

    # title
    s, _ = new(dark=True)
    t = text(s, M, 2.3, 11.5, 2.2, "", 54, font=MEDIUM, color="FFFFFF",
             runs=[("A clear title with a ", False), ("highlight", True)], anchor=MSO_ANCHOR.BOTTOM)
    sub = text(s, M, 4.7, 11.5, 0.6, "Subtitle or event name", 22, color="C9CED6")
    pres = text(s, M, 0.6, 6, 0.9, ["Presenter name", "Company"], 16, color="C9CED6")
    add(id="title", name="Title (dark)", kind="title", aliases=("cover",),
        description="Dark cover with a large title, a subtitle and the presenter at the top.",
        use_when="First slide of a deck.",
        fields=(T("title", t, required=True, highlight=True), T("subtitle", sub), T("presenter", pres)),
        example={"title": "A clear title with a **highlight**", "subtitle": "Subtitle or event name",
                 "presenter": "Presenter name\nCompany"})

    # section
    s, _ = new()
    lab = text(s, M, 2.0, 6, 0.5, "Section 1", 16, font=MEDIUM, color=PALETTE.accent)
    t = text(s, M, 2.6, 11.5, 1.8, "", 48, font=MEDIUM, runs=[("Section ", False), ("title", True)])
    sub = text(s, M, 4.6, 11.5, 0.8, "One short line that frames the section.", 20, color="5B6270")
    add(id="section", name="Section opener", kind="section",
        description="Light slide with a small label, a large title and a short subline.",
        use_when="Start a new section of the talk.",
        fields=(T("label", lab), T("title", t, required=True, highlight=True), T("subtitle", sub)),
        example={"label": "Section 1", "title": "Section **title**",
                 "subtitle": "One short line that frames the section."})

    # statement
    s, _ = new()
    t = text(s, M, 2.2, 11.7, 3.0, "", 44, font=MEDIUM, anchor=MSO_ANCHOR.MIDDLE,
             runs=[("One sentence that carries the ", False), ("key message", True)])
    add(id="statement", name="Statement", kind="statement",
        description="One large sentence with a highlight, vertically centred.",
        use_when="A key message, transition or provocation.",
        fields=(T("title", t, required=True, highlight=True),),
        example={"title": "One sentence that carries the **key message**"})

    # agenda
    s, _ = new()
    t = text(s, M, 0.8, 4.5, 1.0, "Agenda", 44, font=MEDIUM)
    nums, labs = [], []
    for i in range(4):
        y = 1.2 + i * 1.3
        nums.append(text(s, 6.2, y, 1.0, 0.8, f"{i + 1:02d}", 28, font=MEDIUM, color=PALETTE.accent))
        labs.append(text(s, 7.3, y, 5.2, 0.8, ["Opening", "The problem", "The solution", "Next steps"][i], 28))
    add(id="agenda", name="Agenda", kind="agenda", aliases=("steps",),
        description="Title at the left, numbered list of up to four items at the right.",
        use_when="Outline of the talk with up to four short items.",
        fields=(T("title", t, default="Agenda"),),
        items=ItemsSpec(targets={"number": ids(nums), "label": ids(labs)}, min=1, max=4,
                        required=("label",), auto={"number": "index2"}),
        example={"items": [{"label": x} for x in ["Opening", "The problem", "The solution", "Next steps"]]})

    # text
    s, _ = new()
    t = text(s, M, 0.8, 4.8, 2.0, "Title and text", 36, font=MEDIUM)
    body = text(s, 6.2, 0.9, 6.3, 5.0, "A short paragraph that explains one idea. Keep it under sixty words "
                "so the audience can read it while you speak.", 20)
    add(id="text", name="Title and text", kind="text",
        description="Title at the left, one paragraph block at the right.",
        use_when="A single idea explained in a short paragraph.",
        fields=(T("title", t, required=True), T("body", body, required=True, hint="Up to ~60 words.")),
        example={"title": "Title and text", "body": "A short paragraph that explains one idea. Keep it under "
                 "sixty words so the audience can read it while you speak."})

    # rows
    s, _ = new()
    rt, rb = [], []
    for i in range(3):
        y = 0.9 + i * 1.9
        rt.append(text(s, M, y, 4.0, 1.2, ["Point one", "Point two", "Point three"][i], 26, font=MEDIUM))
        rb.append(text(s, 5.4, y, 7.1, 1.4, LOREM, 18))
    add(id="rows-3", name="Three titled rows", kind="rows", aliases=("points",),
        description="Up to three rows. Each row has a title at the left and a paragraph at the right.",
        use_when="Two or three points that each need a paragraph.",
        items=ItemsSpec(targets={"title": ids(rt), "body": ids(rb)}, min=1, max=3, required=("title",)),
        example={"items": [{"title": x, "body": LOREM} for x in ["Point one", "Point two", "Point three"]]})

    # cards
    s, _ = new()
    t = text(s, M, 0.7, 11.7, 1.0, "Three cards", 36, font=MEDIUM)
    cards, ct, cb = [], [], []
    cw = (W - 2 * M - 2 * 0.4) / 3
    for i in range(3):
        x = M + i * (cw + 0.4)
        cards.append(box(s, x, 2.2, cw, 4.2, GREY, MSO_SHAPE.ROUNDED_RECTANGLE))
        ct.append(text(s, x + 0.3, 2.5, cw - 0.6, 1.2, ["First card", "Second card", "Third card"][i], 24,
                       font=MEDIUM))
        cb.append(text(s, x + 0.3, 3.8, cw - 0.6, 2.4, LOREM, 16))
    add(id="cards-3", name="Three cards", kind="cards", aliases=("points",),
        description="Headline and three grey cards, each with a title and a paragraph.",
        use_when="Exactly three points that each need a sentence or two.",
        fields=(T("title", t, required=True),),
        items=ItemsSpec(targets={"title": ids(ct), "body": ids(cb)}, min=3, max=3, required=("title",),
                        extra={"card": ids(cards)},
                        fixed_note="Three cards are part of the design. Provide exactly three items."),
        example={"title": "Three cards", "items": [{"title": x, "body": LOREM}
                                                   for x in ["First card", "Second card", "Third card"]]})

    # text-image
    s, _ = new()
    t = text(s, M, 0.9, 5.2, 1.6, "Text with an image", 34, font=MEDIUM)
    body = text(s, M, 2.7, 5.2, 3.2, LOREM, 18)
    pic = picture(s, 6.6, 0.8, 5.9, 5.6)
    add(id="text-image", name="Text with image", kind="text-image",
        description="Title and paragraph at the left, image at the right.",
        use_when="A point supported by a photo or screenshot.",
        fields=(T("title", t, required=True), T("body", body), I("image", pic, required=True)),
        example={"title": "Text with an image", "body": LOREM, "image": "https://example.com/photo.jpg"})

    # images-2
    s, _ = new()
    pics = [picture(s, M + i * 6.0, 0.8, 5.7, 5.6) for i in range(2)]
    add(id="images-2", name="Two images", kind="image",
        description="Two framed images side by side.",
        use_when="Two photos or screenshots with margins.",
        items=ItemsSpec(targets={"image": ids(pics)}, min=2, max=2, required=("image",),
                        types={"image": "image"}),
        example={"items": [{"image": "https://example.com/photo.jpg"}] * 2})

    # stat
    s, _ = new()
    val = text(s, M, 1.6, 8.0, 3.0, "48%", 160, font=MEDIUM, color=PALETTE.accent, anchor=MSO_ANCHOR.BOTTOM)
    lab = text(s, 8.6, 3.4, 3.9, 1.2, "of teams use the template", 24)
    add(id="big-number", name="Big number", kind="stat",
        description="Very large number at the left and a label at the right.",
        use_when="One key metric.",
        fields=(T("value", val, required=True, hint="e.g. 48%"), T("label", lab, required=True)),
        example={"value": "48%", "label": "of teams use the template"})

    # stats-3
    s, _ = new()
    t = text(s, M, 0.7, 11.7, 1.0, "Three numbers", 36, font=MEDIUM)
    sv, sl = [], []
    for i in range(3):
        x = M + i * 3.95
        sv.append(text(s, x, 2.6, 3.6, 1.6, ["12", "3x", "99%"][i], 72, font=MEDIUM, color=PALETTE.accent))
        sl.append(text(s, x, 4.3, 3.6, 1.0, ["new layouts", "faster decks", "on brand"][i], 20))
    add(id="stats-3", name="Three stats", kind="stats",
        description="Headline, then three accent numbers with a short label under each.",
        use_when="Two or three supporting metrics.",
        fields=(T("title", t, required=True),),
        items=ItemsSpec(targets={"value": ids(sv), "label": ids(sl)}, min=2, max=3, required=("value",)),
        example={"title": "Three numbers", "items": [{"value": v, "label": lb} for v, lb in
                                                     [("12", "new layouts"), ("3x", "faster decks"),
                                                      ("99%", "on brand")]]})

    # bar chart
    s, _ = new()
    t = text(s, M, 0.7, 11.7, 0.9, "Bar chart", 32, font=MEDIUM)
    sub = text(s, M, 1.5, 11.7, 0.6, "Share of teams, in percent", 18, color="5B6270")
    bars, bv = [], []
    max_w, min_w = W - 2 * M, 1.6
    for i, pct in enumerate((35, 80, 55)):
        y = 2.6 + i * 1.35
        w = max(min_w, max_w * pct / 100)
        bars.append(box(s, M, y, w, 1.0, PALETTE.accent))
        bv.append(text(s, M + w - 1.5, y + 0.2, 1.4, 0.6, f"{pct}%", 22, font=MEDIUM, color="FFFFFF",
                       align=PP_ALIGN.RIGHT))
    add(id="bar-chart", name="Bar chart (percentages)", kind="chart",
        description="Two or three horizontal bars sized to their percentage, with the value inside each bar.",
        use_when="Compare two or three percentages.",
        fields=(T("title", t, required=True), T("subtitle", sub)),
        items=ItemsSpec(targets={"value": ids(bv)}, min=2, max=3, required=("value",), extra={"bar": ids(bars)},
                        hints={"value": "Percentage, e.g. 84%. Bar width follows the number."}),
        handler=BarsHandler(max_width_in=max_w, min_width_in=min_w),
        example={"title": "Bar chart", "subtitle": "Share of teams, in percent",
                 "items": [{"value": "35%"}, {"value": "80%"}, {"value": "55%"}]})

    # quote
    s, _ = new()
    text(s, M, 0.4, 2.0, 2.0, "“", 160, font=MEDIUM, color=PALETTE.accent)
    q = text(s, M, 2.2, 11.7, 2.8, "A short quote from a customer or an expert goes here.", 40, font=MEDIUM)
    att = text(s, M, 5.3, 11.7, 0.6, "Name, Role", 18, color="5B6270")
    add(id="quote", name="Quote", kind="quote",
        description="Large accent quote mark, large quote text and the attribution below.",
        use_when="A customer or expert quote.",
        fields=(T("quote", q, required=True, hint="Keep it under ~25 words."), T("attribution", att)),
        example={"quote": "A short quote from a customer or an expert goes here.", "attribution": "Name, Role"})

    # timeline
    s, _ = new()
    t = text(s, M, 0.7, 11.7, 1.0, "Timeline", 36, font=MEDIUM)
    box(s, M, 3.45, W - 2 * M, 0.06, PALETTE.accent_tints[2])
    dots, tl, tb = [], [], []
    step = (W - 2 * M) / 4
    for i in range(4):
        x = M + i * step
        dots.append(box(s, x, 3.3, 0.36, 0.36, PALETTE.accent, MSO_SHAPE.OVAL))
        tl.append(text(s, x, 2.4, step - 0.2, 0.7, ["Q1", "Q2", "Q3", "Q4"][i], 24, font=MEDIUM))
        tb.append(text(s, x, 4.0, step - 0.2, 1.6, "A short milestone.", 16))
    add(id="timeline", name="Timeline", kind="timeline",
        description="Title and up to four milestones on a horizontal line.",
        use_when="Two to four dated milestones.",
        fields=(T("title", t, default="Timeline"),),
        items=ItemsSpec(targets={"label": ids(tl), "body": ids(tb)}, min=2, max=4, required=("label",),
                        extra={"dot": ids(dots)}),
        example={"title": "Timeline", "items": [{"label": q, "body": "A short milestone."}
                                                for q in ["Q1", "Q2", "Q3", "Q4"]]})

    # code
    s, _ = new()
    t = text(s, M, 0.8, 4.4, 1.4, "Code sample", 34, font=MEDIUM)
    sub = text(s, M, 2.3, 4.4, 1.2, "What the code does, in one line.", 18, color="5B6270")
    box(s, 5.6, 0.8, 6.9, 5.7, PALETTE.text, MSO_SHAPE.ROUNDED_RECTANGLE)
    fn = text(s, 5.9, 0.95, 6.3, 0.5, "example.js", 14, font=MONO, color="C9CED6")
    lines = ["function greet(name) {", "  return `Hello, ${name}`;", "}"]
    nums = text(s, 5.9, 1.6, 0.6, 4.7, [f"{i + 1:02d}" for i in range(3)], 14, font=MONO, color="7D8590")
    code = text(s, 6.5, 1.6, 5.8, 4.7, lines, 14, font=MONO, color="FFFFFF")
    add(id="code", name="Code sample", kind="code",
        description="Title and subtitle at the left, dark code window with file name and line numbers at the right.",
        use_when="Show a short code or JSON snippet (up to ~20 lines).",
        fields=(T("title", t, required=True), T("subtitle", sub), T("filename", fn, default="example.js"),
                T("code", code, required=True, hint="Plain text. Newlines become lines; line numbers are added.")),
        handler=CodeHandler(line_numbers=nums.shape_id, max_lines=20),
        example={"title": "Code sample", "subtitle": "What the code does, in one line.", "filename": "example.js",
                 "code": "\n".join(lines)})

    # closing
    s, _ = new(dark=True)
    t = text(s, M, 2.5, W - 2 * M, 2.0, "Thank you", 72, font=MEDIUM, color="FFFFFF", align=PP_ALIGN.CENTER,
             anchor=MSO_ANCHOR.MIDDLE)
    add(id="closing", name="Thank you", kind="closing",
        description="Dark slide with a large centred closing line.",
        use_when="Last slide.",
        fields=(T("title", t, default="Thank you"),),
        example={})

    # master-layout designs (python-pptx default layouts, widened)
    layouts += [
        LayoutSpec(id="title-content", name="Title and content", kind="text", source="layout", ref=ref("content"),
                   description="Title at the top and a text block below.",
                   use_when="A title and several short paragraphs or bullets.",
                   fields=(FieldSpec(name="title", target=0, required=True),
                           FieldSpec(name="body", target=1, required=True)),
                   example={"title": "Title and content", "body": ["First point", "Second point", "Third point"]}),
        LayoutSpec(id="section-header", name="Section header", kind="section", source="layout", ref=ref("section"),
                   description="Plain section header with a title and a short line.",
                   use_when="A quiet section break.",
                   fields=(FieldSpec(name="title", target=0, required=True),
                           FieldSpec(name="subtitle", target=1)),
                   example={"title": "Section header", "subtitle": "A short line"}),
        LayoutSpec(id="image-caption", name="Image with caption", kind="image", source="layout", ref=ref("picture"),
                   description="Title, one large image and a caption.",
                   use_when="One image that needs a title and a caption.",
                   fields=(FieldSpec(name="title", target=0, required=True),
                           FieldSpec(name="image", target=1, type="image", required=True),
                           FieldSpec(name="caption", target=2)),
                   example={"title": "Image with caption", "image": "https://example.com/photo.jpg",
                            "caption": "A short caption"}),
    ]
    assert footer_ref is not None
    return prs, layouts, footer_ref


def main() -> None:
    prs, layouts, foot = build()
    OUT.mkdir(parents=True, exist_ok=True)
    pptx = OUT / TEMPLATE_FILE
    prs.save(pptx)
    missing = [n for n in FONT_FILES if not (OUT / "fonts" / n).exists()]
    if missing:  # the fonts (OFL, see fonts/OFL.txt) are committed with the pack
        raise SystemExit(f"missing fonts in {OUT / 'fonts'}: {missing}")
    example_images(ROOT / "examples" / "assets", ROOT / "src" / "deckwright" / "assets" / "example.jpg")
    pack = Pack(
        id="sample",
        name="Deckwright sample",
        description="Neutral 16:9 sample template: white, dark grey and one blue accent, Manrope.",
        source_sha256=sha256(pptx),
        generator_version="sample-0.1",
        status="confirmed",
        brand=Brand(
            font=FONT,
            replace_font=None,
            diagram_fonts={"Regular": "Manrope-Regular.ttf", "Medium": "ManropeMedium-Regular.ttf",
                           "SemiBold": "Manrope-Bold.ttf"},
            highlight=Highlight(theme_color="accent1", hex=(PALETTE.accent,)),
            colors={"text": PALETTE.text, "white": "FFFFFF", "grey": GREY, "accent": PALETTE.accent},
            palette=PALETTE,
            fonts={"headline": MEDIUM, "body": FONT, "mono": MONO},
            type_scale_pt={"title": 54, "headline": 36, "body": 20, "caption": 16},
            slide_size_in=(W, H),
            rules=("Use one accent colour. Mark a few key words with **double asterisks**.",
                   "Keep card titles under 10 words.",
                   "Vary layouts. Do not put every section on the same text slide."),
        ),
        footer=foot,
        guide="Good rhythm: title -> agenda -> section -> content slides -> closing.",
        layouts=tuple(layouts),
    )
    write_pack(pack, OUT)
    print(f"wrote {OUT} ({len(layouts)} layouts, {len(prs.slides)} slides)")


if __name__ == "__main__":
    main()
