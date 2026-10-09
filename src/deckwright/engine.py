"""Build decks from the template by cloning designed slides or using master layouts."""

from __future__ import annotations

import copy
import io
import json
import math
import re
from dataclasses import dataclass, field
from typing import Any

from lxml import etree
from PIL import Image as PILImage
from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn

from . import diagram
from .images import check_local_path, load_image, placeholder_png
from .pack import LayoutSpec, Template

EMU_PER_IN = 914400
EMU_PER_PT = 12700
SHAPE_TAGS = {qn("p:sp"), qn("p:pic"), qn("p:grpSp"), qn("p:graphicFrame"), qn("p:cxnSp")}
R_NS = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"


@dataclass
class BuildResult:
    data: bytes
    slides: list[dict[str, Any]]
    warnings: list[str] = field(default_factory=list)
    todos: list[dict[str, Any]] = field(default_factory=list)  # placeholder images to replace by hand
    assets: dict[str, bytes] = field(default_factory=dict)  # file name -> bytes (editable .excalidraw files)


class TemplateError(ValueError):
    pass


# --------------------------------------------------------------------------- helpers


def _shape_by_id(root: etree._Element, shape_id: int) -> etree._Element | None:
    for c in root.iter(qn("p:cNvPr")):
        if c.get("id") == str(shape_id):
            el = c.getparent().getparent()
            return el
    return None


def _placeholder_by_idx(root: etree._Element, idx: int) -> etree._Element | None:
    for ph in root.iter(qn("p:ph")):
        if int(ph.get("idx", "0")) == idx:
            return ph.getparent().getparent().getparent()
    return None


def _rel_refs(root: etree._Element):
    """(element, attribute) pairs for every relationship reference (r:embed, r:id, r:link, ...)."""
    for el in root.iter():
        for attr in el.attrib:
            if attr.startswith(R_NS):
                yield el, attr


def _remove(el: etree._Element | None) -> None:
    if el is not None and el.getparent() is not None:
        el.getparent().remove(el)


def _index_label(i: int) -> str:
    return f"{i + 1:02d}"


def _as_paragraphs(value: Any) -> list[str]:
    if value is None:
        return []
    if isinstance(value, (list, tuple)):
        return [str(v) for v in value]
    return [str(value)]


def _paragraph_texts(el: etree._Element) -> list[str]:
    """Text per visual line: paragraphs split at line breaks."""
    out: list[str] = []
    for p in el.iter(qn("a:p")):
        cur = ""
        for c in p:
            if c.tag == qn("a:r"):
                cur += "".join(t.text or "" for t in c.iter(qn("a:t")))
            elif c.tag == qn("a:br"):
                out.append(cur)
                cur = ""
        out.append(cur)
    return [t for t in out if t.strip()] or []


def _percent(value: Any) -> float:
    """First number in a value. '1,200' and '1.200,5' use thousands separators; '12,5' is a decimal."""
    m = re.search(r"-?\d[\d.,]*", str(value))
    if not m:
        return 0.0
    num = m.group(0).rstrip(".,")
    if re.fullmatch(r"-?\d{1,3}(,\d{3})+(\.\d+)?", num):
        num = num.replace(",", "")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+(,\d+)?", num):
        num = num.replace(".", "").replace(",", ".")
    else:
        num = num.replace(",", ".")
    try:
        return float(num)
    except ValueError:
        return 0.0


# --------------------------------------------------------------------------- text


def _rpr_is_highlight(rpr: etree._Element | None, theme_color: str, hex_set: set[str]) -> bool:
    if rpr is None:
        return False
    fill = rpr.find(qn("a:solidFill"))
    if fill is None:
        return False
    scheme = fill.find(qn("a:schemeClr"))
    srgb = fill.find(qn("a:srgbClr"))
    return (scheme is not None and scheme.get("val") == theme_color) or (
        srgb is not None and srgb.get("val", "").upper() in hex_set
    )


def _make_highlight(rpr: etree._Element, theme_color: str) -> etree._Element:
    rpr = copy.deepcopy(rpr)
    for tag in ("a:noFill", "a:solidFill", "a:gradFill", "a:blipFill", "a:pattFill", "a:grpFill"):
        for c in rpr.findall(qn(tag)):
            rpr.remove(c)
    fill = etree.SubElement(rpr, qn("a:solidFill"))
    etree.SubElement(fill, qn("a:schemeClr"), val=theme_color)
    rpr.remove(fill)
    ln = rpr.find(qn("a:ln"))
    rpr.insert(0 if ln is None else list(rpr).index(ln) + 1, fill)
    return rpr


class _TextSlot:
    """A text shape plus the formatting captured from its template content."""

    def __init__(self, sp: etree._Element, theme_color: str, hex_set: set[str]):
        self.sp = sp
        self.theme_color = theme_color
        self.body = sp.find(qn("p:txBody"))
        if self.body is None:
            raise TemplateError("shape has no text body")
        self.templates: list[tuple[Any, Any, Any]] = []
        self.pink: etree._Element | None = None
        for p in self.body.findall(qn("a:p")):
            runs = p.findall(qn("a:r"))
            ppr = p.find(qn("a:pPr"))
            end = p.find(qn("a:endParaRPr"))
            rpr = runs[0].find(qn("a:rPr")) if runs else end
            for r in runs:
                r_rpr = r.find(qn("a:rPr"))
                if self.pink is None and _rpr_is_highlight(r_rpr, theme_color, hex_set):
                    self.pink = r_rpr
            if runs or not self.templates:
                self.templates.append((ppr, rpr, end))
        if not self.templates:
            self.templates.append((None, None, None))

    def fill(self, paragraphs: list[tuple[int, str]], highlight: bool, plain_ppr: bool) -> None:
        for p in self.body.findall(qn("a:p")):
            self.body.remove(p)
        for tpl_i, text in paragraphs:
            ppr, rpr, end = self.templates[min(tpl_i, len(self.templates) - 1)]
            base = copy.deepcopy(rpr) if rpr is not None else etree.Element(qn("a:rPr"))
            base.tag = qn("a:rPr")
            base.set("lang", base.get("lang", "en"))
            pink = copy.deepcopy(self.pink) if self.pink is not None else _make_highlight(base, self.theme_color)
            p = etree.SubElement(self.body, qn("a:p"))
            if ppr is not None:
                p.append(copy.deepcopy(ppr))
            elif plain_ppr:
                new_ppr = etree.SubElement(p, qn("a:pPr"), marL="0", indent="0")
                etree.SubElement(new_ppr, qn("a:buNone"))
            for li, line in enumerate(text.split("\n")):
                if li:
                    br = etree.SubElement(p, qn("a:br"))
                    br.append(copy.deepcopy(base))
                parts = line.split("**") if highlight else [line.replace("**", "")]
                for si, seg in enumerate(parts):
                    if not seg:
                        continue
                    r = etree.SubElement(p, qn("a:r"))
                    r.append(copy.deepcopy(pink if si % 2 else base))
                    t = etree.SubElement(r, qn("a:t"))
                    t.text = seg
            if end is not None:
                e = copy.deepcopy(end)
                e.tag = qn("a:endParaRPr")
                p.append(e)


# --------------------------------------------------------------------------- builder


class DeckBuilder:
    def __init__(self, template: Template, allow_local_files: bool = True):
        self.template = template
        self.brand = template.pack.brand
        self.style = diagram.Style.from_template(template)
        self.prs = Presentation(str(template.pptx))
        self.allow_local_files = allow_local_files
        self._orig = list(self.prs.slides)
        self._orig_ids = list(self.prs.slides._sldIdLst)
        self.warnings: list[str] = []
        self.report: list[dict[str, Any]] = []
        self.todos: list[dict[str, Any]] = []
        self.assets: dict[str, bytes] = {}

    def _placeholder_png(self) -> bytes:
        return placeholder_png(self.brand.palette.accent_tints[1])

    # ---- slide creation

    def _src_slide(self, number: int):
        if not 1 <= number <= len(self._orig):
            raise TemplateError(f"template has no slide {number}")
        return self._orig[number - 1]

    def _copy_rels(self, src_part, dst_part, roots: list[etree._Element]) -> None:
        """Relate dst_part to the targets that the copied elements reference, and remap their rIds."""
        refs = [ref for root in roots for ref in _rel_refs(root)]
        rid_map: dict[str, str] = {}
        for rid in {el.get(attr) for el, attr in refs}:
            rel = src_part.rels.get(rid)
            if rel is None or rel.reltype in (RT.SLIDE_LAYOUT, RT.NOTES_SLIDE):
                continue
            if rel.is_external:
                rid_map[rid] = dst_part.relate_to(rel.target_ref, rel.reltype, is_external=True)
            else:
                rid_map[rid] = dst_part.relate_to(rel.target_part, rel.reltype)
        for el, attr in refs:
            v = el.get(attr)
            if v in rid_map:
                el.set(attr, rid_map[v])

    @staticmethod
    def _prune_rels(slide) -> None:
        """Drop relationships no longer referenced by the slide XML (replaced or removed images)."""
        used = {el.get(attr) for el, attr in _rel_refs(slide._element)}
        for rid, rel in list(slide.part.rels.items()):
            if rel.reltype in (RT.SLIDE_LAYOUT, RT.NOTES_SLIDE):
                continue
            if rid not in used:
                slide.part.rels.pop(rid)

    def clone_slide(self, number: int):
        src = self._src_slide(number)
        new = self.prs.slides.add_slide(src.slide_layout)
        tree = new.shapes._spTree
        for el in list(tree):
            if el.tag in SHAPE_TAGS:
                tree.remove(el)
        src_csld = src._element.find(qn("p:cSld"))
        dst_csld = new._element.find(qn("p:cSld"))
        copies = [copy.deepcopy(el) for el in src.shapes._spTree if el.tag in SHAPE_TAGS]
        bg = src_csld.find(qn("p:bg"))
        bg = copy.deepcopy(bg) if bg is not None else None
        self._copy_rels(src.part, new.part, copies + ([bg] if bg is not None else []))
        if bg is not None:
            dst_csld.insert(0, bg)
        for c in copies:
            tree.append(c)
        return new

    def layout_slide(self, index: int, footer: bool = True):
        layouts = self.prs.slide_layouts
        if not 0 <= index < len(layouts):
            raise TemplateError(f"template has no layout {index}")
        new = self.prs.slides.add_slide(layouts[index])
        pack_footer = self.template.pack.footer
        if footer and pack_footer is not None:
            src = self._src_slide(pack_footer.slide)
            el = _shape_by_id(src.shapes._spTree, pack_footer.shape)
            if el is not None:
                c = copy.deepcopy(el)
                self._copy_rels(src.part, new.part, [c])
                new.shapes._spTree.append(c)
        return new

    # ---- filling

    def _warn(self, slide_no: int, msg: str) -> None:
        self.warnings.append(f"slide {slide_no}: {msg}")

    def _frame_in(self, el: etree._Element, slide) -> tuple[float, float]:
        """Size of an image frame in inches (from the shape, or its layout placeholder)."""
        sp_pr = el.find(qn("p:spPr"))
        xfrm = sp_pr.find(qn("a:xfrm")) if sp_pr is not None else None
        if xfrm is None:
            layout_ph = self._layout_ph(el, slide)
            xfrm = layout_ph.find(qn("p:spPr")).find(qn("a:xfrm")) if layout_ph is not None else None
        if xfrm is None:
            return 16.0, 9.0
        ext = xfrm.find(qn("a:ext"))
        return int(ext.get("cx")) / EMU_PER_IN, int(ext.get("cy")) / EMU_PER_IN

    @staticmethod
    def _placeholder_label(source: Any) -> str | None:
        if isinstance(source, dict) and source.get("placeholder") and not source.get("excalidraw"):
            return str(source["placeholder"])
        if isinstance(source, str) and source.lower().startswith("placeholder:"):
            return source.split(":", 1)[1].strip() or "Image"
        return None

    def _image_data(self, source: Any, slide_no: int, name: str, slide, el: etree._Element) -> bytes:
        """Bytes for an image slot. Sources: URL, data URI, path, 'placeholder: label',
        {"placeholder": label}, {"excalidraw": scene} or a path to a .excalidraw file."""
        w_in, h_in = self._frame_in(el, slide)
        if not source:
            self._warn(slide_no, f"'{name}' has no image; a pink placeholder is used")
            return self._placeholder_png()
        label = self._placeholder_label(source)
        if label is not None:
            self.todos.append({"slide": slide_no, "field": name, "label": label,
                               "size_in": [round(w_in, 2), round(h_in, 2)]})
            return diagram.placeholder(label, w_in, h_in, self.style)
        if isinstance(source, dict) and "excalidraw" not in source:
            self._warn(slide_no, f"'{name}' image source has keys {sorted(source)}; use a URL, data URI, path, "
                                 "'placeholder: ...', {\"placeholder\": ...} or {\"excalidraw\": ...}. "
                                 "A pink placeholder is used")
            return self._placeholder_png()
        scene = source.get("excalidraw") if isinstance(source, dict) else None
        if scene is None and isinstance(source, str) and source.endswith(".excalidraw"):
            try:
                scene = json.loads(check_local_path(source, self.allow_local_files).read_text())
            except (OSError, ValueError) as exc:  # PermissionError is an OSError
                self._warn(slide_no, f"'{name}' diagram could not be loaded ({exc}); a pink placeholder is used")
                return self._placeholder_png()
        if scene is not None:
            try:
                scene = diagram.normalise(scene)
                png = diagram.render(scene, w_in, h_in, self.style)
                safe = re.sub(r"[^a-z0-9]+", "-", name.lower()).strip("-")
                self.assets[f"slide{slide_no:02d}-{safe}.excalidraw"] = json.dumps(scene, indent=1).encode()
                return png
            except Exception as exc:  # noqa: BLE001
                self._warn(slide_no, f"'{name}' diagram could not be rendered ({exc}); a pink placeholder is used")
                return self._placeholder_png()
        try:
            return load_image(source, allow_local=self.allow_local_files)
        except Exception as exc:  # noqa: BLE001
            self._warn(slide_no, f"'{name}' image could not be loaded ({exc}); a pink placeholder is used")
            return self._placeholder_png()

    def _set_picture(self, slide, el: etree._Element, data: bytes) -> None:
        image_part, rid = slide.part.get_or_add_image_part(io.BytesIO(data))
        with PILImage.open(io.BytesIO(data)) as im:
            iw, ih = im.size
        if el.tag == qn("p:pic"):
            blip_fill = el.find(qn("p:blipFill"))
        else:  # empty picture placeholder (p:sp): convert via python-pptx
            ph = next(s for s in slide.placeholders if s._element is el)
            ph.insert_picture(io.BytesIO(data))
            return
        blip = blip_fill.find(qn("a:blip"))
        blip.set(qn("r:embed"), rid)
        ext = el.find(qn("p:spPr")).find(qn("a:xfrm")).find(qn("a:ext"))
        fw, fh = int(ext.get("cx")), int(ext.get("cy"))
        for c in blip_fill.findall(qn("a:srcRect")):
            blip_fill.remove(c)
        rect = etree.Element(qn("a:srcRect"))
        img_ratio, frame_ratio = iw / ih, fw / fh
        if img_ratio > frame_ratio:
            crop = int((1 - frame_ratio / img_ratio) / 2 * 100000)
            rect.set("l", str(crop))
            rect.set("r", str(crop))
        elif img_ratio < frame_ratio:
            crop = int((1 - img_ratio / frame_ratio) / 2 * 100000)
            rect.set("t", str(crop))
            rect.set("b", str(crop))
        blip_fill.insert(list(blip_fill).index(blip) + 1, rect)
        for c in blip_fill.findall(qn("a:tile")):
            blip_fill.remove(c)
        if blip_fill.find(qn("a:stretch")) is None:
            st = etree.SubElement(blip_fill, qn("a:stretch"))
            etree.SubElement(st, qn("a:fillRect"))

    def _fill_text(self, slide_no: int, el: etree._Element, name: str,
                   paragraphs: list[tuple[int, str]], highlight: bool, plain_ppr: bool, fit: bool) -> None:
        original = _paragraph_texts(el)
        hl = self.brand.highlight
        slot = _TextSlot(el, hl.theme_color, {h.upper() for h in hl.hex})
        slot.fill(paragraphs, highlight, plain_ppr)
        self._enforce_brand_font(el)
        if fit:
            self._shrink_to_fit(slide_no, el, name, original)

    # ---- brand font

    def _inherited_latin(self, el: etree._Element) -> str | None:
        ph = el.find(".//" + qn("p:ph"))
        if ph is None or ph.get("idx") == "4294967295":
            return None  # text boxes inherit the theme font (Arial)
        match = self._layout_ph(el, self._current)
        if match is not None:
            for lat in match.iter(qn("a:latin")):
                return lat.get("typeface")
        master = self._current.slide_layout.slide_master._element
        style = "p:titleStyle" if ph.get("type") in ("title", "ctrTitle") else "p:bodyStyle"
        lat = master.find(f".//{qn(style)}/{qn('a:lvl1pPr')}/{qn('a:defRPr')}/{qn('a:latin')}")
        return lat.get("typeface") if lat is not None else None

    def _enforce_brand_font(self, el: etree._Element) -> None:
        """Some slots may be left on a theme font the template sets by mistake (``replace_font``).
        The brand allows only its own font, so set it where nothing else is set."""
        bad_font = self.brand.replace_font
        rprs = list(el.iter(qn("a:rPr"), qn("a:endParaRPr")))
        missing = [r for r in rprs if r.find(qn("a:latin")) is None]
        if not missing:
            return
        found = next((lat.get("typeface") for r in rprs for lat in r.findall(qn("a:latin"))), None)
        if found and found != bad_font:
            font = found  # the designer set a font on part of the box: use it throughout
        else:
            inherited = self._inherited_latin(el)
            if inherited and not inherited.startswith("+") and inherited != bad_font:
                return
            font = self.brand.font
        rprs = missing
        after = {qn(t) for t in ("a:ea", "a:cs", "a:sym", "a:hlinkClick", "a:hlinkMouseOver", "a:rtl", "a:extLst")}
        for r in rprs:
            pos = next((i for i, c in enumerate(r) if c.tag in after), len(r))
            for j, tag in enumerate(("a:latin", "a:ea", "a:cs", "a:sym")):
                if r.find(qn(tag)) is None:
                    r.insert(pos + j, etree.Element(qn(tag), typeface=font))

    # ---- text fitting (estimate, then shrink)

    def _effective_size(self, el: etree._Element, slide) -> float:
        sizes = [int(r.get("sz")) for r in el.iter(qn("a:rPr"), qn("a:endParaRPr"), qn("a:defRPr")) if r.get("sz")]
        if sizes:
            return max(sizes) / 100
        ph = el.find(".//" + qn("p:ph"))
        if ph is not None:
            layout_el = slide.slide_layout._element
            match = _placeholder_by_idx(layout_el, int(ph.get("idx", "0")))
            if match is not None:
                sizes = [int(r.get("sz")) for r in match.iter(qn("a:rPr"), qn("a:defRPr"), qn("a:endParaRPr")) if r.get("sz")]
                if sizes:
                    return max(sizes) / 100
            master = slide.slide_layout.slide_master._element
            style = "p:titleStyle" if ph.get("type") in ("title", "ctrTitle") else "p:bodyStyle"
            lvl1 = master.find(f".//{qn(style)}/{qn('a:lvl1pPr')}/{qn('a:defRPr')}")
            if lvl1 is not None and lvl1.get("sz"):
                return int(lvl1.get("sz")) / 100
        return 18.0

    def _shrink_to_fit(self, slide_no: int, el: etree._Element, name: str, original: list[str]) -> None:
        """Estimate wrapped line count; shrink the font if the text is taller than the box.

        The template's own sample text always counts as fitting, which
        calibrates the estimate against the designer's box sizes.
        """
        slide = self._current
        xfrm = el.find(qn("p:spPr")).find(qn("a:xfrm"))
        if xfrm is None:
            layout_ph = self._layout_ph(el, slide)
            xfrm = layout_ph.find(qn("p:spPr")).find(qn("a:xfrm")) if layout_ph is not None else None
        if xfrm is None:
            return
        ext = xfrm.find(qn("a:ext"))
        body_pr = el.find(qn("p:txBody")).find(qn("a:bodyPr"))

        def ins(key: str, default: int) -> int:
            return int(body_pr.get(key, default)) if body_pr is not None else default

        w_pt = (int(ext.get("cx")) - ins("lIns", 91440) - ins("rIns", 91440)) / EMU_PER_PT
        h_pt = (int(ext.get("cy")) - ins("tIns", 45720) - ins("bIns", 45720)) / EMU_PER_PT
        if w_pt <= 0:
            return
        grows = body_pr is not None and body_pr.find(qn("a:spAutoFit")) is not None
        size = self._effective_size(el, slide)
        texts = _paragraph_texts(el)

        def lines(sz: float, paras: list[str]) -> int:
            cpl = max(1.0, w_pt / (sz * 0.5))
            return sum(max(1, math.ceil(len(t) / cpl)) for t in paras)

        def fits(sz: float) -> bool:
            n = lines(sz, texts)
            if n <= 1 or n <= lines(size, original):
                return True
            return n * sz * 1.12 <= h_pt * (1.5 if grows else 1.0)

        if fits(size):
            return
        scale = 1.0
        while scale > 0.55 and not fits(size * scale):
            scale -= 0.05
        for r in el.iter(qn("a:rPr"), qn("a:endParaRPr")):
            own = int(r.get("sz")) / 100 if r.get("sz") else size
            r.set("sz", str(int(own * scale * 100)))
        if not fits(size * scale):
            self._warn(slide_no, f"'{name}' text is too long for its box even at {int(scale * 100)}% size; shorten it")
        else:
            self._warn(slide_no, f"'{name}' text was shrunk to {int(scale * 100)}% to fit")

    def _layout_ph(self, el: etree._Element, slide):
        ph = el.find(".//" + qn("p:ph"))
        if ph is None:
            return None
        return _placeholder_by_idx(slide.slide_layout._element, int(ph.get("idx", "0")))

    # ---- main entry per slide

    def add(self, layout: LayoutSpec, content: dict[str, Any], notes: str | None = None,
            fit: bool = True, footer: bool = True) -> None:
        slide_no = len(self.report) + 1
        if layout.source == "slide":
            slide = self.clone_slide(layout.ref)
        else:
            slide = self.layout_slide(layout.ref, footer and layout.footer)
        self._current = slide
        root = slide.shapes._spTree
        by_target = (lambda t: _shape_by_id(root, t)) if layout.source == "slide" else (lambda t: _placeholder_by_idx(root, t))
        plain = layout.source == "layout"
        used = set()

        for sid in layout.strip:
            _remove(by_target(sid))

        # scalar fields, grouped by target (several fields may share one shape)
        groups: dict[int, list] = {}
        for f in layout.fields:
            groups.setdefault(f.target, []).append(f)
        for target, fields in groups.items():
            el = by_target(target)
            if el is None:
                raise TemplateError(f"{layout.id}: target {target} not found")
            first = fields[0]
            if first.type == "image":
                val = content.get(first.name)
                used.add(first.name)
                if val:
                    self._set_picture(slide, el, self._image_data(val, slide_no, first.name, slide, el))
                elif first.required:
                    self._set_picture(slide, el, self._image_data(None, slide_no, first.name, slide, el))
                else:
                    _remove(el)
                    for extra in first.also_remove:
                        _remove(by_target(extra))
                continue
            paragraphs: list[tuple[int, str]] = []
            for f in fields:
                used.add(f.name)
                val = content.get(f.name)
                if val in (None, "", []):
                    val = f.default
                if val in (None, "", []):
                    if f.required:
                        self._warn(slide_no, f"required field '{f.name}' is empty")
                    continue
                if layout.handler is not None and layout.handler.type == "code" and f.name == "code":
                    continue
                tpl = f.paragraph or 0
                for i, text in enumerate(_as_paragraphs(val)):
                    paragraphs.append((tpl + i, text))
            if layout.handler is not None and layout.handler.type == "code" and first.name == "code":
                self._fill_code(slide_no, root, el, str(content.get("code", "")), fit, layout.handler)
                continue
            if not paragraphs:
                _remove(el)
                for f in fields:
                    for extra in f.also_remove:
                        _remove(by_target(extra))
                continue
            if first.width_in:
                self._set_width(el, first.width_in)
            self._fill_text(slide_no, el, first.name, paragraphs, any(f.highlight for f in fields), plain, fit)

        # repeated items
        if layout.items:
            self._fill_items(slide, slide_no, layout, content.get("items") or [], by_target, plain, fit)
            used.add("items")

        dropped = sorted(k for k in content if k not in used)
        if dropped:
            self._warn(slide_no, f"layout '{layout.id}' has no slot for: {', '.join(dropped)}")

        self._prune_rels(slide)
        self._write_notes(slide, slide_no, notes)
        self.report.append({"slide": slide_no, "layout": layout.id})

    def _fill_items(self, slide, slide_no: int, layout: LayoutSpec, items: list[dict[str, Any]],
                    by_target, plain: bool, fit: bool) -> None:
        spec = layout.items
        assert spec is not None
        if len(items) < spec.min:
            self._warn(slide_no, f"'{layout.id}' expects at least {spec.min} items, got {len(items)}")
        if len(items) > spec.max:
            self._warn(slide_no, f"'{layout.id}' holds {spec.max} items; {len(items) - spec.max} dropped")
            items = items[: spec.max]
        items = [i if isinstance(i, dict) else {spec.fields[0]: i} for i in items]
        for pos in range(spec.max):
            item = items[pos] if pos < len(items) else None
            for fname, targets in spec.targets.items():
                el = by_target(targets[pos])
                if el is None:
                    continue
                is_image = spec.types.get(fname) == "image"
                if item is None:
                    if is_image and pos < spec.min:  # fixed image slot: keep the frame
                        self._set_picture(slide, el, self._placeholder_png())
                    else:
                        _remove(el)
                    continue
                if is_image:
                    self._set_picture(slide, el, self._image_data(item.get(fname), slide_no,
                                                                 f"items[{pos}].{fname}", slide, el))
                    continue
                if fname in spec.auto:
                    val = _index_label(pos)
                else:
                    val = item.get(fname)
                if val in (None, "", []):
                    if fname in spec.required and item:
                        self._warn(slide_no, f"items[{pos}].{fname} is empty")
                    _remove(el)
                    continue
                self._fill_text(slide_no, el, f"items[{pos}].{fname}", [(0, str(v)) for v in _as_paragraphs(val)],
                                False, plain, fit)
            for _, targets in spec.extra.items():
                if item is None:
                    _remove(by_target(targets[pos]))
        if layout.handler is not None and layout.handler.type == "bars":
            self._layout_bars(slide, items, spec, by_target, layout.handler)
        elif layout.handler is not None and layout.handler.type == "columns":
            self._layout_columns(slide, items, spec, by_target, layout.handler)

    def _set_width(self, el: etree._Element, width_in: float) -> None:
        sp_pr = el.find(qn("p:spPr"))
        xfrm = sp_pr.find(qn("a:xfrm"))
        if xfrm is None:  # placeholder inherits its box from the layout: copy it
            layout_ph = self._layout_ph(el, self._current)
            src = layout_ph.find(qn("p:spPr")).find(qn("a:xfrm")) if layout_ph is not None else None
            if src is None:
                return
            xfrm = copy.deepcopy(src)
            sp_pr.insert(0, xfrm)
        xfrm.find(qn("a:ext")).set("cx", str(int(width_in * EMU_PER_IN)))

    # ---- handlers

    def _xfrm(self, el):
        x = el.find(qn("p:spPr")).find(qn("a:xfrm"))
        return x.find(qn("a:off")), x.find(qn("a:ext"))

    def _layout_bars(self, slide, items, spec, by_target, handler) -> None:
        max_w = int(handler.max_width_in * EMU_PER_IN)
        min_w = int(handler.min_width_in * EMU_PER_IN)
        for pos, item in enumerate(items[: spec.max]):
            bar, label = by_target(spec.extra["bar"][pos]), by_target(spec.targets["value"][pos])
            if bar is None or label is None:
                continue
            pct = max(0.0, min(100.0, float(item.get("percent", _percent(item.get("value"))))))
            b_off, b_ext = self._xfrm(bar)
            w = max(min_w, int(max_w * pct / 100))
            b_ext.set("cx", str(w))
            l_off, l_ext = self._xfrm(label)
            l_off.set("x", str(int(b_off.get("x")) + w - int(l_ext.get("cx"))))

    def _layout_columns(self, slide, items, spec, by_target, handler) -> None:
        bottom = int(handler.baseline_in * EMU_PER_IN)
        max_h = int(handler.max_height_in * EMU_PER_IN)
        min_h = int(handler.min_height_in * EMU_PER_IN)
        for pos, item in enumerate(items[: spec.max]):
            col, value = by_target(spec.extra["column"][pos]), by_target(spec.targets["value"][pos])
            if col is None:
                continue
            pct = max(0.0, min(100.0, float(item.get("percent", _percent(item.get("value"))))))
            h = max(min_h, int(max_h * pct / 100))
            c_off, c_ext = self._xfrm(col)
            c_ext.set("cy", str(h))
            c_off.set("y", str(bottom - h))
            if value is not None:
                v_off, _ = self._xfrm(value)
                v_off.set("y", str(bottom - h))

    def _fill_code(self, slide_no: int, root, el, code: str, fit: bool, handler) -> None:
        lines = code.rstrip("\n").split("\n") if code else [""]
        if len(lines) > handler.max_lines:
            self._warn(slide_no, f"code has {len(lines)} lines; about {handler.max_lines} fit")
        self._fill_text(slide_no, el, "code", [(0, ln if ln else " ") for ln in lines], False, False, fit)
        numbers = _shape_by_id(root, handler.line_numbers) if handler.line_numbers is not None else None
        if numbers is not None:
            self._fill_text(slide_no, numbers, "line numbers", [(0, _index_label(i)) for i in range(len(lines))],
                            False, False, fit)
            if fit:  # keep line numbers the same size as the code
                sz = next((r.get("sz") for r in el.iter(qn("a:rPr")) if r.get("sz")), None)
                if sz:
                    for r in numbers.iter(qn("a:rPr"), qn("a:endParaRPr")):
                        r.set("sz", sz)

    # ---- raw layout access

    def add_raw(self, index: int, placeholders: dict[str, Any], notes: str | None, footer: bool, fit: bool) -> None:
        slide_no = len(self.report) + 1
        slide = self.layout_slide(index, footer)
        self._current = slide
        root = slide.shapes._spTree
        filled = set()
        for key, val in placeholders.items():
            try:
                idx = int(key)
            except (TypeError, ValueError):
                raise TemplateError(f"placeholder key '{key}' is not a placeholder idx number") from None
            el = _placeholder_by_idx(root, idx)
            if el is None:
                self._warn(slide_no, f"layout {index} has no placeholder idx {idx}")
                continue
            filled.add(idx)
            ph_type = el.find(".//" + qn("p:ph")).get("type")
            if ph_type == "pic":
                self._set_picture(slide, el, self._image_data(val, slide_no, f"placeholder {idx}", slide, el))
            else:
                self._fill_text(slide_no, el, f"placeholder {idx}", [(0, t) for t in _as_paragraphs(val)], True, True, fit)
        for ph in list(root.iter(qn("p:ph"))):
            if int(ph.get("idx", "0")) not in filled:
                _remove(ph.getparent().getparent().getparent())
        self._prune_rels(slide)
        self._write_notes(slide, slide_no, notes)
        self.report.append({"slide": slide_no, "layout_index": index, "layout_name": self.prs.slide_layouts[index].name})

    def _write_notes(self, slide, slide_no: int, notes: str | None) -> None:
        todo = [f"TODO: replace the placeholder '{t['field']}' with: {t['label']}"
                for t in self.todos if t["slide"] == slide_no]
        text = "\n".join([*( [notes] if notes else []), *todo])
        if text:
            slide.notes_slide.notes_text_frame.text = text

    # ---- output

    def _register_notes_master(self) -> None:
        """python-pptx adds a notes master when the template has none, but does not list it in
        presentation.xml. PowerPoint accepts that; Keynote refuses the file as invalid."""
        pres = self.prs.part._element
        if pres.find(qn("p:notesMasterIdLst")) is not None:
            return
        rid = next((r.rId for r in self.prs.part.rels.values() if r.reltype == RT.NOTES_MASTER), None)
        if rid is None:
            return
        lst = etree.Element(qn("p:notesMasterIdLst"))
        etree.SubElement(lst, qn("p:notesMasterId")).set(qn("r:id"), rid)
        pres.find(qn("p:sldMasterIdLst")).addnext(lst)  # the schema order puts it right after

    def build(self, title: str | None = None, author: str | None = None) -> BuildResult:
        lst = self.prs.slides._sldIdLst
        for sld in self._orig_ids:
            rid = sld.rId
            lst.remove(sld)
            self.prs.part.drop_rel(rid)
        self._orig_ids = []
        self._register_notes_master()
        if title:
            self.prs.core_properties.title = title
        if author:
            self.prs.core_properties.author = author
        buf = io.BytesIO()
        self.prs.save(buf)
        return BuildResult(buf.getvalue(), self.report, self.warnings, self.todos, self.assets)


def layout_or_raise(template: Template, layout_id: str) -> LayoutSpec:
    try:
        return template.by_id[layout_id]
    except KeyError:
        raise TemplateError(f"unknown layout '{layout_id}'") from None


def template_layouts(template: Template) -> list[dict[str, Any]]:
    """Describe all master layouts with their placeholders (raw access)."""
    prs = Presentation(str(template.pptx))
    out = []
    for i, lay in enumerate(prs.slide_layouts):
        phs = []
        for ph in sorted(lay.placeholders, key=lambda s: (round(s.top / EMU_PER_IN), s.left)):
            f = ph.placeholder_format
            phs.append({
                "idx": f.idx,
                "type": str(f.type).split(".")[-1].split(" ")[0].lower(),
                "box_in": [round(v / EMU_PER_IN, 2) for v in (ph.left, ph.top, ph.width, ph.height)],
            })
        out.append({"index": i, "name": lay.name, "placeholders": phs})
    return out
