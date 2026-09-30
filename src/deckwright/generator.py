"""Generate a template pack from a .pptx by inspecting its shapes.

The generator reads every visible sample slide and the master layouts that the
samples do not cover. It finds text, image and item slots, guesses a content
kind for each design, reads the brand from the theme, and returns a draft
``Pack``. It needs no network and no LLM. Results it is unsure about go into
``Pack.issues`` and a low ``LayoutSpec.confidence``, for review.
"""

from __future__ import annotations

import hashlib
import re
from collections import Counter
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from lxml import etree
from pptx import Presentation
from pptx.oxml.ns import qn

from .pack import (
    EXAMPLE_IMAGE,
    BarsHandler,
    Brand,
    CodeHandler,
    ColumnsHandler,
    FieldSpec,
    Footer,
    Highlight,
    Issue,
    ItemsSpec,
    LayoutSpec,
    Pack,
    Palette,
    sha256,
)

GENERATOR_VERSION = "0.1"
EMU = 914400
TOL = 0.12  # inches: alignment tolerance
NON_CONTENT_PH = {"dt", "ftr", "sldNum", "hdr", "sldImg"}
MONO_HINTS = ("mono", "code", "courier", "consol", "menlo")
QUOTE_MARKS = set("\"'“”‘’«»„")
NUMBER_RE = re.compile(r"^[\s$€£¥+~<>-]*\d[\d.,]*\s*(%|x|k|m|b|bn|h|hrs?|min|s|pp|\+|m\+|k\+)?\s*$", re.I)
DATE_RE = re.compile(r"\b(19|20)\d{2}\b|\b(jan|feb|mar|apr|may|jun|jul|aug|sep|oct|nov|dec)[a-z]*\b", re.I)
INDEX_RE = re.compile(r"^0?\d{1,2}\.?$")
FILLER = ("lorem", "ipsum", "here", "title", "headline", "subtitle", "subheadline", "text", "description",
          "placeholder", "name", "caption", "label", "your ")
RULE_START = ("use ", "do ", "don't", "do not", "keep ", "avoid ", "always ", "never ", "only ", "make ", "place ",
              "every ", "make sure", "limit ")


# --------------------------------------------------------------------------- inspection


@dataclass
class Shape:
    id: int
    el: Any
    kind: str  # text | picture | auto | group
    x: float
    y: float
    w: float
    h: float
    text: str = ""
    size: float = 18.0
    font: str | None = None
    mono: bool = False
    highlight: tuple[str | None, str | None] | None = None  # (scheme slot, srgb) of an accent run
    font_explicit: bool = False  # set on the shape or its layout placeholder, not only by the theme
    ph_type: str | None = None
    ph_idx: int | None = None
    image_sha: str | None = None
    decoration_text: bool = False

    @property
    def cx(self) -> float:
        return self.x + self.w / 2

    @property
    def cy(self) -> float:
        return self.y + self.h / 2

    @property
    def lines(self) -> list[str]:
        return [ln for ln in re.split(r"[\n\x0b]", self.text) if ln.strip()]

    @property
    def words(self) -> int:
        return len(self.text.split())

    def contains(self, other: Shape, pad: float = 0.05) -> bool:
        return (self.x - pad <= other.x and other.x + other.w <= self.x + self.w + pad
                and self.y - pad <= other.y and other.y + other.h <= self.y + self.h + pad)

    def overlaps(self, other: Shape) -> bool:
        ix = min(self.x + self.w, other.x + other.w) - max(self.x, other.x)
        iy = min(self.y + self.h, other.y + other.h) - max(self.y, other.y)
        return ix > 0.05 and iy > 0.05 and ix * iy > 0.5 * min(self.w * self.h, other.w * other.h)


def _sizes(el) -> list[float]:
    return [int(r.get("sz")) / 100 for r in el.iter(qn("a:rPr"), qn("a:endParaRPr"), qn("a:defRPr")) if r.get("sz")]


def _latin(el) -> list[str]:
    return [lat.get("typeface") for lat in el.iter(qn("a:latin")) if lat.get("typeface")]


class _Inspector:
    def __init__(self, prs):
        self.prs = prs
        self.master = prs.slide_master
        self.theme_fonts = self._theme_fonts()

    def _theme_fonts(self) -> dict[str, str]:
        from pptx.opc.constants import RELATIONSHIP_TYPE as RT

        theme = etree.fromstring(self.master.part.part_related_by(RT.THEME).blob)
        out = {}
        for key, tag in (("major", "a:majorFont"), ("minor", "a:minorFont")):
            lat = theme.find(f".//{qn(tag)}/{qn('a:latin')}")
            if lat is not None:
                out[key] = lat.get("typeface")
        return out

    def theme_colors(self) -> dict[str, str]:
        from pptx.opc.constants import RELATIONSHIP_TYPE as RT

        theme = etree.fromstring(self.master.part.part_related_by(RT.THEME).blob)
        scheme = theme.find(f".//{qn('a:clrScheme')}")
        out = {}
        for c in scheme if scheme is not None else []:
            name = etree.QName(c).localname
            srgb, sys_ = c.find(qn("a:srgbClr")), c.find(qn("a:sysClr"))
            val = srgb.get("val") if srgb is not None else (sys_.get("lastClr") if sys_ is not None else None)
            if val:
                out[name] = val.upper()
        return out

    def _base_placeholder(self, el, layout):
        ph = el.find(".//" + qn("p:ph"))
        if ph is None or layout is None:
            return None
        idx = int(ph.get("idx", "0"))
        for s in layout.placeholders:
            if s.placeholder_format.idx == idx:
                return s._element
        return None

    def _master_style(self, ph_type: str | None) -> tuple[float | None, str | None]:
        style = "p:titleStyle" if ph_type in ("title", "ctrTitle") else "p:bodyStyle"
        d = self.master._element.find(f".//{qn(style)}/{qn('a:lvl1pPr')}/{qn('a:defRPr')}")
        if d is None:
            return None, None
        lat = d.find(qn("a:latin"))
        return (int(d.get("sz")) / 100 if d.get("sz") else None), (lat.get("typeface") if lat is not None else None)

    def _resolve_font(self, name: str | None, ph_type: str | None) -> str | None:
        if name in ("+mj-lt", None) and ph_type in ("title", "ctrTitle"):
            return self.theme_fonts.get("major") if name else None
        if name == "+mj-lt":
            return self.theme_fonts.get("major")
        if name == "+mn-lt":
            return self.theme_fonts.get("minor")
        return name

    def shapes(self, container, layout=None) -> list[Shape]:
        """Top-level shapes of a slide or layout, with effective text size and font."""
        out = []
        for sh in container.shapes:
            el = sh._element
            try:
                x, y, w, h = (v / EMU for v in (sh.left, sh.top, sh.width, sh.height))
            except TypeError:
                continue
            ph = el.find(".//" + qn("p:ph"))
            ph_type = ph.get("type", "body") if ph is not None else None
            ph_idx = int(ph.get("idx", "0")) if ph is not None else None
            if ph_type in NON_CONTENT_PH:
                continue
            tag = etree.QName(el).localname
            if tag == "grpSp":
                kind = "group"
            elif tag == "pic" or ph_type == "pic":
                kind = "picture"
            elif sh.has_text_frame and (sh.text_frame.text.strip() or ph is not None):
                kind = "text"
            elif tag in ("sp", "cxnSp"):
                kind = "auto"
            else:
                continue
            s = Shape(sh.shape_id, el, kind, x, y, w, h, ph_type=ph_type, ph_idx=ph_idx)
            if kind == "picture" and tag == "pic":
                blip = el.find(".//" + qn("a:blip"))
                rid = blip.get(qn("r:embed")) if blip is not None else None
                try:
                    s.image_sha = hashlib.sha1(container.part.related_part(rid).blob).hexdigest() if rid else None
                except KeyError:
                    s.image_sha = None
            if kind == "text":
                s.text = sh.text_frame.text
                base = self._base_placeholder(el, layout)
                sizes = _sizes(el) or (_sizes(base) if base is not None else [])
                fonts = _latin(el) or (_latin(base) if base is not None else [])
                m_size, m_font = self._master_style(ph_type) if ph is not None else (None, None)
                s.size = max(sizes) if sizes else (m_size or 18.0)
                font = fonts[0] if fonts else (m_font or ("+mj-lt" if ph_type in ("title", "ctrTitle") else "+mn-lt"))
                s.font = self._resolve_font(font, ph_type)
                s.font_explicit = bool(fonts or m_font) and not str(font).startswith("+")
                s.mono = any(h in (s.font or "").lower() for h in MONO_HINTS)
                s.highlight = self._highlight(el)
                t = s.text.strip()
                if t and (set(t) <= QUOTE_MARKS or t in ("‹#›", "<#>")):
                    s.decoration_text = True
            out.append(s)
        return out

    @staticmethod
    def _run_color(r) -> tuple[str | None, str | None]:
        rpr = r.find(qn("a:rPr"))
        fill = rpr.find(qn("a:solidFill")) if rpr is not None else None
        if fill is None:
            return None, None
        sc, rgb = fill.find(qn("a:schemeClr")), fill.find(qn("a:srgbClr"))
        return (sc.get("val") if sc is not None else None), (rgb.get("val").upper() if rgb is not None else None)

    def _highlight(self, el) -> tuple[str | None, str | None] | None:
        """The colour of a run that differs from the other runs in the same shape."""
        runs = [r for r in el.iter(qn("a:r")) if "".join(t.text or "" for t in r.iter(qn("a:t"))).strip()]
        if len(runs) < 2:
            return None
        colors = [self._run_color(r) for r in runs]
        counts = Counter(colors)
        if len(counts) < 2:
            return None
        low = min(counts.values())
        rare = [c for c in counts if counts[c] == low and c != (None, None)]
        if not rare:
            return None
        # ties: a theme accent beats a literal colour, and white or black are rarely the highlight
        return min(rare, key=lambda c: (not (c[0] or "").startswith("accent"), c[1] in ("FFFFFF", "000000")))


# --------------------------------------------------------------------------- series (items)


@dataclass
class Series:
    members: list[Shape]
    kind: str

    @property
    def n(self) -> int:
        return len(self.members)


def _similar(a: float, b: float, rel: float = 0.08) -> bool:
    return abs(a - b) <= max(TOL, rel * max(a, b))


def _row_major(shapes: list[Shape]) -> list[Shape]:
    rows: list[list[Shape]] = []
    for s in sorted(shapes, key=lambda s: s.y):
        for row in rows:
            if abs(row[0].y - s.y) < 0.4:
                row.append(s)
                break
        else:
            rows.append([s])
    return [s for row in rows for s in sorted(row, key=lambda s: s.x)]


def _find_series(shapes: list[Shape]) -> list[Series]:
    """Group shapes of the same kind, style and size that line up in a row, column or grid."""
    parent = {s.id: s.id for s in shapes}

    def find(i):
        while parent[i] != i:
            parent[i] = parent[parent[i]]
            i = parent[i]
        return i

    def link(a, b):
        parent[find(a.id)] = find(b.id)

    for i, a in enumerate(shapes):
        for b in shapes[i + 1:]:
            if a.kind != b.kind or a.kind == "group":
                continue
            if a.kind == "text" and (abs(a.size - b.size) > 1 or a.mono != b.mono):
                continue
            same_w, same_h = _similar(a.w, b.w), _similar(a.h, b.h)
            row = abs(a.y - b.y) < TOL or abs(a.cy - b.cy) < TOL or abs(a.y + a.h - b.y - b.h) < TOL
            col = abs(a.x - b.x) < TOL or abs(a.cx - b.cx) < TOL
            if a.kind == "text":
                ok = (row and same_w) or (col and (same_h or (same_w and a.font == b.font)))
            elif a.kind == "picture":
                ok = row or col
            else:
                ok = (row and (same_w or same_h)) or (col and (same_h or same_w))
            if ok:
                link(a, b)
    groups: dict[int, list[Shape]] = {}
    for s in shapes:
        groups.setdefault(find(s.id), []).append(s)
    return [Series(_row_major(g), g[0].kind) for g in groups.values() if len(g) >= 2]


def _pair(a: Series, b: Series) -> Series | None:
    """``b`` reordered to match ``a`` member by member, if every member of ``a`` has its own neighbour in ``b``."""
    if a.n != b.n:
        return None
    left = list(b.members)
    out = []
    for x in a.members:
        near = [y for y in left if abs(x.cx - y.cx) < max(x.w, y.w) / 2 or abs(x.cy - y.cy) < max(x.h, y.h) / 2 + 0.3]
        if not near:
            return None
        y = min(near, key=lambda y: (x.cx - y.cx) ** 2 + (x.cy - y.cy) ** 2)
        left.remove(y)
        out.append(y)
    return Series(out, b.kind)


# --------------------------------------------------------------------------- slot naming


def _is_number(t: str) -> bool:
    t = t.strip()
    return bool(t) and len(t) <= 12 and len(t.split()) <= 2 and bool(NUMBER_RE.match(t))


def _is_filler(t: str) -> bool:
    low = t.lower()
    return any(f in low for f in FILLER)


@dataclass
class Draft:
    fields: list[FieldSpec] = field(default_factory=list)
    items: dict[str, list[int]] = field(default_factory=dict)
    item_types: dict[str, str] = field(default_factory=dict)
    extra: dict[str, list[int]] = field(default_factory=dict)
    auto: dict[str, str] = field(default_factory=dict)
    example: dict[str, Any] = field(default_factory=dict)
    item_examples: list[dict[str, Any]] = field(default_factory=list)
    series_orient: str = "row"
    handler: Any = None
    has_line: bool = False
    has_dots: bool = False
    bar_like: str | None = None
    unmapped: list[int] = field(default_factory=list)
    texts: list[Shape] = field(default_factory=list)
    quote_mark: bool = False
    cards: bool = False
    cover: bool = False  # a large photo behind or beside a large title
    title_text: str = ""


def _example_text(s: Shape) -> str | list[str]:
    lines = s.lines
    return lines[0] if len(lines) == 1 else "\n".join(lines) if len(lines) <= 3 else lines


class _Mapper:
    def __init__(self, slide_w: float, slide_h: float, is_layout: bool):
        self.W, self.H = slide_w, slide_h
        self.is_layout = is_layout

    def map(self, shapes: list[Shape], context: list[Shape] = ()) -> Draft:
        """``context`` holds master-layout artwork behind the slide: never slots, only hints."""
        d = Draft()
        texts = [s for s in shapes if s.kind == "text" and not s.decoration_text]
        pics = [s for s in shapes if s.kind == "picture"]
        autos = [s for s in shapes if s.kind == "auto"]
        d.texts = texts
        d.quote_mark = any(s.decoration_text and set(s.text.strip()) <= QUOTE_MARKS for s in [*shapes, *context])
        lines = [a for a in [*autos, *context] if a.kind == "auto"]
        d.has_line = any(a.h < 0.15 and a.w > self.W * 0.5 or a.w < 0.15 and a.h > self.H * 0.5 for a in lines)
        self.context = [c for c in context if c.kind in ("auto", "picture")]
        used: set[int] = set()

        # --- code window: mono multi-line text, optional line-number box beside it
        mono = [s for s in texts if s.mono]
        code = max((s for s in mono if not all(INDEX_RE.match(ln.strip()) for ln in s.lines)),
                   key=lambda s: len(s.lines) * s.w, default=None)
        if code is not None and len(code.lines) >= 2:
            nums = next((s for s in texts if s is not code and len(s.lines) >= 2 and s.lines and all(INDEX_RE.match(ln.strip())
                                                                             for ln in s.lines)), None)
            fname = next((s for s in mono if s not in (code, nums) and len(s.lines) == 1), None)
            d.fields.append(FieldSpec(name="code", target=code.id, required=True,
                                      hint="Plain text. Newlines become lines; line numbers are added."))
            d.example["code"] = "\n".join(code.lines)
            used.add(code.id)
            if fname is not None:
                d.fields.append(FieldSpec(name="filename", target=fname.id, default=fname.text.strip()))
                used.add(fname.id)
            max_lines = max(5, int(code.h * 72 / (code.size * 1.2)))
            d.handler = CodeHandler(line_numbers=nums.id if nums else None, max_lines=max_lines)
            if nums is not None:
                used.add(nums.id)

        # --- item series
        cand = [s for s in texts + pics if s.id not in used]
        series = [x for x in _find_series(cand)]
        auto_series = _find_series(autos)
        # bars / columns: value texts inside a series of rectangles of varying length
        for a in auto_series:
            inside = [[t for t in texts if t.id not in used and m.overlaps(t) and m.contains(t, pad=0.3)]
                      for m in a.members]
            if all(len(x) >= 1 for x in inside) and all(_is_number(x[0].text) for x in inside):
                widths = {round(m.w, 1) for m in a.members}
                heights = {round(m.h, 1) for m in a.members}
                d.bar_like = "bars" if len(widths) > 1 and len(heights) == 1 else "columns" if len(heights) > 1 else None
                if d.bar_like:
                    vals = [x[0] for x in inside]
                    by = (lambda i: (a.members[i].y, a.members[i].x)) if d.bar_like == "bars" else \
                        (lambda i: (a.members[i].x, a.members[i].y))
                    order = sorted(range(a.n), key=by)
                    members = [a.members[i] for i in order]
                    vals = [vals[i] for i in order]
                    d.items["value"] = [v.id for v in vals]
                    d.extra["bar" if d.bar_like == "bars" else "column"] = [m.id for m in members]
                    d.item_examples = [{"value": v.text.strip()} for v in vals]
                    used |= {v.id for v in vals}
                    if d.bar_like == "bars":
                        left = min(m.x for m in members)
                        d.handler = BarsHandler(max_width_in=round(self.W - 2 * left, 2),
                                                min_width_in=round(min(v.w for v in vals) + 0.2, 2))
                    else:
                        base = max(m.y + m.h for m in members)
                        tallest = max(members, key=lambda m: m.h)
                        d.handler = ColumnsHandler(baseline_in=round(base, 2),
                                                   max_height_in=round(min(base - 0.5, tallest.h * 1.1), 2),
                                                   min_height_in=round(min(m.h for m in members), 2))
                    labels = [t for t in texts if t.id not in used and any(m.contains(t, 0.3) for m in members)]
                    lab_series = [s for s in _find_series(labels) if s.n == len(members)]
                    if lab_series:
                        d.items["label"] = [t.id for t in _row_major(lab_series[0].members)]
                        for ex, t in zip(d.item_examples, _row_major(lab_series[0].members)):
                            ex["label"] = _example_text(t)
                        used |= set(d.items["label"])
                    series = [s for s in _find_series([s for s in cand if s.id not in used])]
                    break

        if not d.bar_like:
            series = [s for s in series if not any(m.id in used for m in s.members)]
            self._items(d, series, autos, used)

        # --- scalar fields
        rest = [s for s in texts if s.id not in used]
        free_pics = [p for p in pics if p.id not in used]
        self._scalars(d, rest, free_pics, used)
        d.unmapped = [s.id for s in texts if s.id not in used]
        return d

    def _items(self, d: Draft, series: list[Series], autos: list[Shape], used: set[int]) -> None:
        if not series:
            return
        # the biggest family of mutually paired series becomes the item set
        best: list[Series] = []
        for s in sorted(series, key=lambda s: (-s.n, s.kind != "picture", -s.members[0].size)):
            fam = [s] + [q for q in (_pair(s, o) for o in series if o is not s) if q is not None]
            if len(fam) > len(best) or (len(fam) == len(best) and fam[0].n > (best[0].n if best else 0)):
                best = fam
        n = best[0].n
        if n < 2:
            return
        text_series = [s for s in best if s.kind == "text"]
        pic_series = [s for s in best if s.kind == "picture"]
        # order text series by role: index numbers, values, then by font size (title before body)
        names: dict[int, str] = {}
        for s in text_series:
            samples = [m.text.strip() for m in s.members]
            if all(INDEX_RE.match(t) for t in samples if t) and any(samples):
                names[id(s)] = "number"
            elif any(samples) and all(_is_number(t) for t in samples if t):
                names[id(s)] = "value"
        empty = [s for s in text_series if id(s) not in names and not any(m.text.strip() for m in s.members)]
        if len(text_series) >= 3 and empty:
            narrow = min(empty, key=lambda s: s.members[0].w)
            if narrow.members[0].w < 1.6 and narrow.members[0].h < 1.2:
                names[id(narrow)] = "number"
        rest = [s for s in text_series if id(s) not in names]
        if all(not m.text.strip() for s in rest for m in s.members):
            # empty boxes: a title-type placeholder, else the upper box, is the title
            rest.sort(key=lambda s: (s.members[0].ph_type == "body", s.members[0].y, s.members[0].h))
        else:
            rest.sort(key=lambda s: (-s.members[0].size, s.members[0].y))
        has_value = "value" in names.values()
        role = ["label"] if has_value else []
        if len(rest) == 1:
            s = rest[0]
            long_text = max(m.words for m in s.members) > 12 or (s.members[0].h > 1.5 and s.members[0].size < 30)
            small = s.members[0].size < 20 or s.members[0].h < 0.5
            if has_value or "number" in names.values():
                names[id(s)] = "label"
            elif pic_series and small:
                names[id(s)] = "caption"
            else:
                names[id(s)] = "body" if long_text and pic_series else "title"
        else:
            order = role + ["title", "body", "caption", "label"]
            for s, name in zip(rest, [o for o in order]):
                names[id(s)] = name
        for s in text_series:
            name = names[id(s)]
            if name in d.items:
                continue
            d.items[name] = [m.id for m in s.members]
            if name == "number":
                d.auto["number"] = "index2"
            used |= {m.id for m in s.members}
        for i, s in enumerate(pic_series[:1]):
            d.items["image"] = [m.id for m in s.members]
            d.item_types["image"] = "image"
            used |= {m.id for m in s.members}
        # containers and dots that repeat with the items
        first = [d.items[k] for k in d.items][0]
        members = {m.id: m for s in best for m in s.members}
        anchors = [members[i] for i in first]
        for a in _find_series(autos) + _find_series(self.context):
            if a.n != n:
                continue
            ordered = _row_major(a.members)
            in_layout = any(c in self.context for c in ordered)
            if all(c.x <= t.cx <= c.x + c.w and c.y <= t.cy <= c.y + c.h for c, t in zip(ordered, anchors)):
                if in_layout:
                    d.cards = True
                else:
                    d.extra["card"] = [c.id for c in ordered]
                    d.cards = True
            elif all(max(c.w, c.h) < 0.8 for c in ordered):
                if not in_layout:
                    d.extra["dot"] = [c.id for c in ordered]
                d.has_dots = True
        d.series_orient = "col" if len({round(a.y, 1) for a in anchors}) == n else "row"
        d.item_examples = []
        for pos in range(n):
            ex = {}
            for name, ids in d.items.items():
                if name in d.auto:
                    continue
                if d.item_types.get(name) == "image":
                    ex[name] = EXAMPLE_IMAGE
                    continue
                t = members[ids[pos]].text.strip()
                if t:
                    ex[name] = _example_text(members[ids[pos]])
            d.item_examples.append(ex)

    def _scalars(self, d: Draft, texts: list[Shape], pics: list[Shape], used: set[int]) -> None:
        names = {f.name for f in d.fields}

        def add(name, s, **kw):
            if name in names:
                return
            names.add(name)
            used.add(s.id)
            d.fields.append(FieldSpec(name=name, target=s.id, **kw))
            if s.kind == "picture":
                d.example[name] = EXAMPLE_IMAGE
            elif s.text.strip():
                d.example[name] = _example_text(s)

        for p in sorted(pics, key=lambda p: -p.w * p.h)[:2] if not isinstance(d.handler, CodeHandler) else []:
            if p.w * p.h < 0.5:  # tiny pictures (logos, icons) stay as artwork
                continue
            add("image" if "image" not in names else "image2", p, type="image", required=True)

        values = [s for s in texts if _is_number(s.text)]
        others = [s for s in texts if s not in values]
        if values and not d.items:
            v = max(values, key=lambda s: s.size)
            add("value", v, required=True)
            texts = [s for s in texts if s is not v]
            others = [s for s in texts if not _is_number(s.text)]
        quote = next((s for s in others if s.text.strip()[:1] in QUOTE_MARKS or s.text.strip()[-1:] in QUOTE_MARKS),
                     None)
        if not others:
            return
        by_size = sorted(others, key=lambda s: (-s.size, -s.w * s.h))
        if self.is_layout:  # no sample text: trust the placeholder type
            by_size.sort(key=lambda s: (s.ph_type not in ("title", "ctrTitle"), -s.size))
        title = by_size[0]
        near = [p for p in pics if p.w * p.h >= 0.5]
        if near and len(by_size) == 1 and (title.size < 20 or title.h < 0.6) and title.words <= 12:
            add("caption", title)
            return
        if "value" in names:
            # stat: the text next to the number is its label
            add("label", title, required=True)
            for s in by_size[1:2]:
                add("caption", s)
            return
        if quote is not None or d.quote_mark:
            q = quote or title
            add("quote", q, required=True, hint="Keep it under ~25 words.")
            for s in [s for s in by_size if s is not q][:1]:
                add("attribution", s)
            return
        sample = title.text.strip()
        d.title_text = sample
        fixed = sample and not _is_filler(sample) and len(sample.split()) <= 3
        # a small label above the title (e.g. "Section 1")
        for s in sorted(by_size[1:], key=lambda s: s.x):
            top = (s.cy < title.cy and s.y < self.H * 0.35 and s.size < title.size * 0.7 and s.w < self.W * 0.5
                   and s.h < 1.0)
            if not top or s.words > 6:
                continue
            if DATE_RE.search(s.text) and "date" not in names:
                add("date", s, hint="e.g. 13 October 2026")
            elif "date" in names or len(s.lines) > 1 or "," in s.text:
                add("presenter", s, hint="Name and company. Use \\n for a line break.")
            elif s.words <= 4 and "label" not in names:
                add("label", s)
        width = None
        if title.w < self.W * 0.4 and not any(o is not title and o.y < title.y + title.h and o.y + o.h > title.y
                                              and o.x > title.x + title.w for o in texts + pics):
            width = round(self.W * 0.7, 2)
        d.cover = title.size >= 60 and any(p.w * p.h >= 0.35 * self.W * self.H for p in pics)
        add("title", title, required=not fixed, highlight=title.highlight is not None,
            default=sample if fixed else None, width_in=width)
        rest = [s for s in by_size[1:] if s.id not in used]
        for s in rest:
            near_pic = any(abs(s.x - p.x) < 0.5 and 0 <= s.y - (p.y + p.h) < 1.0 for p in pics)
            if near_pic and s.words <= 12:
                add("caption", s)
            elif s.words >= 8 or len(s.lines) > 2 or (not s.text.strip() and s.h > 1.5):
                add("body" if "body" not in names else "body2", s)
            elif "subtitle" not in names:
                add("subtitle", s)
            elif "body" not in names:
                add("body", s)


# --------------------------------------------------------------------------- kinds


def _kind(d: Draft, pos: int, total: int, texts_all: str, n_pics: int) -> tuple[str, tuple[str, ...], float]:
    """(kind, aliases, confidence) from the slot pattern."""
    f = {x.name for x in d.fields}
    it = set(d.items)
    n = len(next(iter(d.items.values()))) if d.items else 0
    low = texts_all.lower()
    tl = d.title_text.lower()
    if re.search(r"thank|follow me|contact|q\s*&\s*a|questions", tl) and not it:
        return ("qa", ("closing",), 0.75) if re.search(r"q\s*&\s*a|questions", tl) else ("closing", (), 0.75)
    if isinstance(d.handler, CodeHandler):
        return "code", (), 0.95
    if d.bar_like:
        return "chart", (), 0.9
    if "quote" in f:
        return "quote", (), 0.85
    if "number" in it:
        if "step" in d.title_text.lower():
            return "steps", ("agenda",), 0.8
        return "agenda", (), 0.85
    if "value" in it:
        return "stats", (), 0.85
    if "value" in f:
        return "stat", (), 0.85
    if it and d.cards and "image" not in it:
        return "cards", ("points",), 0.8
    if it and d.has_dots or (it and d.has_line and n >= 3 and "image" not in it):
        return "timeline", (), 0.75
    if it == {"image"} or it == {"image", "caption"}:
        if {"title", "body"} <= f:
            return "text-image", (), 0.8
        return "image", (), 0.85
    if "image" in it:
        return "text-image", ("points",), 0.7
    if it:
        if d.cards:
            return "cards", ("points",), 0.8
        if d.series_orient == "col":
            return "rows", ("points",), 0.7
        return "columns", ("points",), 0.65
    if "image" in f:
        others = f - {"image", "image2", "caption"}
        if d.cover and "body" not in f and "title" in f:
            return "title", ("cover",), 0.7
        if not others:
            return "image", (), 0.8
        if pos == 0 or "date" in f or "presenter" in f:
            return "title", ("cover",), 0.6
        return ("text-image", (), 0.75) if "body" in f else ("image", (), 0.65)
    if re.search(r"\bq\s*&\s*a\b|questions", low):
        return "qa", ("closing",), 0.7
    if re.search(r"thank|follow me|contact", low) or (pos == total - 1 and total > 3):
        return "closing", (), 0.7
    if "agenda" in low:
        return "agenda", (), 0.5
    if not f:
        return "divider", (), 0.6
    if "date" in f or "presenter" in f or (pos == 0 and len(f) >= 2):
        return "title", ("cover",), 0.75
    if "label" in f:
        return "section", (), 0.75
    if f <= {"title"}:
        return "statement", (), 0.7
    if "body" in f:
        return "text", (), 0.7
    if f <= {"title", "subtitle", "presenter"}:
        return "statement", (), 0.55
    return "text", (), 0.4


# --------------------------------------------------------------------------- brand


NAME_HINTS = [
    (r"title slide|cover", ("title", ("cover",), 0.8)),
    (r"section", ("section", (), 0.8)),
    (r"quote", ("quote", (), 0.8)),
    (r"agenda", ("agenda", (), 0.8)),
    (r"timeline", ("timeline", (), 0.8)),
    (r"thank|closing|end", ("closing", (), 0.8)),
]


def _name_hint(name: str, d: Draft) -> tuple[str, tuple[str, ...], float] | None:
    """Kind from a meaningful layout name. Exported Google Slides names carry none."""
    low = name.lower()
    for pat, hint in NAME_HINTS:
        if re.search(pat, low):
            return hint
    return None


DEFAULT_FONTS = {"Arial", "Calibri", "Helvetica", "Helvetica Neue", "Times New Roman", "Cambria", "Aptos",
                 "Open Sans", "Roboto"}
WEIGHTS = {"thin", "extralight", "light", "regular", "medium", "semibold", "bold", "extrabold", "black", "book",
           "demibold", "heavy", "italic"}


def _family(font: str | None) -> str:
    """'Brand Sans SemiBold' -> 'Brand Sans'."""
    parts = (font or "").split()
    while len(parts) > 1 and parts[-1].lower() in WEIGHTS:
        parts.pop()
    return " ".join(parts)


def _brand(prs, ins: _Inspector, design: list[list[Shape]], hidden_text: list[str],
           highlights: list[tuple[str | None, str | None]]) -> Brand:
    colors = ins.theme_colors()

    def lum(h: str) -> float:
        r, g, b = (int(h[i:i + 2], 16) for i in (0, 2, 4))
        return 0.2126 * r + 0.7152 * g + 0.0722 * b

    dk, lt = colors.get("dk1", "000000"), colors.get("lt1", "FFFFFF")
    text, bg = (dk, lt) if lum(dk) < lum(lt) else (lt, dk)
    acc = [colors.get(f"accent{i}", "888888") for i in range(1, 7)]
    fonts = Counter(s.font for shapes in design for s in shapes
                    if s.kind == "text" and s.font and not s.mono and s.font_explicit)
    mono = Counter(s.font for shapes in design for s in shapes if s.kind == "text" and s.mono and s.font)
    own = Counter({f: n for f, n in fonts.items() if _family(f) not in DEFAULT_FONTS})
    body = (own or fonts).most_common(1)[0][0] if fonts else (ins.theme_fonts.get("minor") or "Arial")
    family = _family(body)
    big = Counter(s.font for shapes in design for s in shapes
                  if s.kind == "text" and s.size >= 36 and s.font and s.font_explicit)
    theme_minor = ins.theme_fonts.get("minor")
    replace = theme_minor if theme_minor and _family(theme_minor) != family else None
    hl = Counter(highlights).most_common(1)
    slot, hexv = hl[0][0] if hl else ("accent1", None)
    if slot is None and hexv:
        slot = next((k for k, v in colors.items() if v == hexv), "accent1")
    hex_set = tuple(sorted({h for s, h in highlights if h} | ({colors[slot]} if slot in colors else set())))
    sizes = sorted({round(s.size) for shapes in design for s in shapes if s.kind == "text"}, reverse=True)
    rules = tuple(t for t in hidden_text if t.lower().startswith(RULE_START) and len(t.split()) <= 40)[:15]
    return Brand(
        font=family,
        replace_font=replace,
        highlight=Highlight(theme_color=slot or "accent1", hex=hex_set),
        colors={k: v for k, v in colors.items() if k.startswith(("dk", "lt", "accent"))},
        palette=Palette(text=text, background=bg, accent=acc[0], accent_tints=(acc[1], acc[2], acc[3]),
                        accent_dark=acc[4], accent_darkest=acc[5]),
        fonts={"body": body, **({"headline": big.most_common(1)[0][0]} if big else {}),
               **({"mono": mono.most_common(1)[0][0]} if mono else {}),
               "families": sorted({f for f in [*fonts, *mono] if _family(f) not in DEFAULT_FONTS})},
        type_scale_pt={f"size_{i + 1}": float(s) for i, s in enumerate(sizes[:8])},
        slide_size_in=(round(prs.slide_width / EMU, 3), round(prs.slide_height / EMU, 3)),
        rules=rules,
    )


# --------------------------------------------------------------------------- main


def _slug(kind: str, n: int | None, taken: set[str]) -> str:
    base = f"{kind}-{n}" if n else kind
    out, i = base, 2
    while out in taken:
        out, i = f"{base}-{chr(ord('a') + i - 2)}" if n else f"{base}-{i}", i + 1
    taken.add(out)
    return out


DESCRIPTIONS = {
    "title": ("Cover slide.", "First slide of a deck."),
    "section": ("Section opener with a label and a large title.", "Start a new section."),
    "statement": ("One large statement.", "A key message or transition."),
    "agenda": ("Numbered list of short items.", "Outline of the talk."),
    "steps": ("Numbered list of steps.", "Next steps or a short procedure."),
    "text": ("Title and a text block.", "A single idea explained in a short paragraph."),
    "rows": ("Stacked rows of titled text.", "A few points that each need a sentence."),
    "columns": ("Side-by-side columns of text.", "Parallel points."),
    "cards": ("Cards with a title and optional text.", "A fixed number of short takeaways."),
    "text-image": ("Text with one or more images.", "A point supported by images."),
    "image": ("Image-led slide.", "Photos or screenshots."),
    "stat": ("One large number with a label.", "One key metric."),
    "stats": ("Several numbers with labels.", "Two or more supporting metrics."),
    "chart": ("Shapes sized to percentages.", "Compare a few percentages."),
    "quote": ("A quote with attribution.", "A customer or expert quote."),
    "timeline": ("Milestones along a line.", "Dated milestones or phases."),
    "code": ("Code window with line numbers.", "A short code or JSON snippet."),
    "closing": ("Closing slide.", "Last slide."),
    "qa": ("Question time slide.", "Q&A."),
    "divider": ("Artwork only, no text.", "A visual pause. Use sparingly."),
}


SAMPLE_TEXT = {
    "title": "Title", "subtitle": "A short subtitle", "body": "One or two short sentences that support the point.",
    "body2": "A second short paragraph.", "label": "Label", "value": "42%", "caption": "A short caption",
    "date": "1 January 2027", "presenter": "Name Surname\nCompany", "quote": "A short quote that makes the point.",
    "attribution": "Name, Role", "code": "const answer = 42;\nconsole.log(answer);", "filename": "example.js",
}


WORDS = ("Short supporting text that explains the point in a few clear words for the audience "
         "and keeps the slide easy to read").split()


def items_has(d: Draft, name: str) -> bool:
    return name in d.items


def _fit_text(name: str, shape: Shape | None, i: int | None = None) -> str:
    """Example text for a slot, cut to what its box holds (same estimate as the engine)."""
    base = SAMPLE_TEXT.get(name, "Text") if not (name == "title" and i is not None) else f"Item {i + 1}"
    if shape is None or name in ("value", "date", "code", "filename", "presenter"):
        return base
    cpl = max(1.0, (shape.w - 0.2) * 72 / (shape.size * 0.5))
    lines = max(1, int((shape.h - 0.1) * 72 / (shape.size * 1.12)))
    cap = int(cpl * lines * 0.75)
    if len(base) <= cap:
        return base
    out = ""
    for w in (base.split() if name in ("title", "label", "subtitle", "caption") else WORDS):
        if len(out) + len(w) + 1 > cap:
            break
        out = f"{out} {w}".strip()
    return out or base.split()[0]


def _fill_examples(fields: list[FieldSpec], items: ItemsSpec | None, example: dict[str, Any],
                   shapes: dict[int, Shape]) -> dict[str, Any]:
    """Neutral example text for slots whose template box has no sample text (master layouts)."""
    ex = dict(example)
    for f in fields:
        if f.name not in ex and not f.default:
            ex[f.name] = EXAMPLE_IMAGE if f.type == "image" else _fit_text(f.name, shapes.get(f.target))
    if items:
        rows = list(ex.get("items") or [])
        rows += [{} for _ in range(items.min - len(rows))] if len(rows) < items.min else []
        if not rows:
            rows = [{} for _ in range(items.max)]
        for i, row in enumerate(rows):
            for name in items.fields:
                if name not in row:
                    if items.types.get(name) == "image":
                        row[name] = EXAMPLE_IMAGE
                    elif name == "value":
                        row[name] = f"{(i + 1) * 20}%"
                    else:
                        row[name] = _fit_text(name, shapes.get(items.targets[name][i]), i)
        ex["items"] = rows
    return ex


@dataclass
class _Candidate:
    source: str
    ref: int
    shapes: list[Shape]
    name: str
    context: list[Shape] = field(default_factory=list)


def generate(pptx: str | Path, pack_id: str, name: str | None = None) -> Pack:
    """Inspect ``pptx`` and return a draft pack."""
    pptx = Path(pptx)
    prs = Presentation(str(pptx))
    ins = _Inspector(prs)
    W, H = prs.slide_width / EMU, prs.slide_height / EMU
    slides = list(prs.slides)
    visible = [(i + 1, s) for i, s in enumerate(slides) if s._element.get("show") != "0"]
    hidden_text = [p.text.strip() for s in slides if s._element.get("show") == "0"
                   for sh in s.shapes if sh.has_text_frame for p in sh.text_frame.paragraphs if p.text.strip()]
    layouts = list(prs.slide_layouts)
    issues: list[Issue] = []

    slide_shapes = {num: ins.shapes(s, s.slide_layout) for num, s in visible}

    # footer: a picture repeated at the same place on most design slides
    seen: Counter = Counter()
    first: dict[tuple, tuple[int, int]] = {}
    for num, shapes in slide_shapes.items():
        for s in shapes:
            if s.kind == "picture" and s.image_sha:
                key = (s.image_sha, round(s.x, 1), round(s.y, 1), round(s.w, 1))
                seen[key] += 1
                first.setdefault(key, (num, s.id))
    footer = None
    footer_keys = {k for k, c in seen.items() if c >= max(2, 0.4 * len(visible))}
    if footer_keys:
        key = max(footer_keys, key=lambda k: seen[k])
        footer = Footer(slide=first[key][0], shape=first[key][1])

    def strip_footer(shapes):
        return [s for s in shapes if not (s.kind == "picture" and s.image_sha and
                                          (s.image_sha, round(s.x, 1), round(s.y, 1), round(s.w, 1)) in footer_keys)]

    cands: list[_Candidate] = []
    used_layouts: dict[int, list[set[int]]] = {}
    for num, s in visible:
        shapes = strip_footer(slide_shapes[num])
        ctx = [x for x in ins.shapes(s.slide_layout) if x.ph_idx is None]
        cands.append(_Candidate("slide", num, shapes, f"Slide {num}", ctx))
        li = layouts.index(s.slide_layout)
        used_layouts.setdefault(li, []).append({sh.ph_idx for sh in shapes if sh.ph_idx is not None})
    for li, lay in enumerate(layouts):
        shapes = ins.shapes(lay, None)
        phs = {s.ph_idx for s in shapes if s.ph_idx is not None}
        deco = [s for s in shapes if s.ph_idx is None and (s.kind in ("picture", "auto", "group") or
                                                            (s.kind == "text" and not s.decoration_text))]
        if li in used_layouts and all(phs <= u for u in used_layouts[li]):
            continue  # a sample slide already shows this design
        if len(phs) == 1 or (not phs and not deco):
            continue
        for sh in shapes:
            if sh.ph_idx is not None:
                sh.id = sh.ph_idx  # master-layout fields address placeholder idx
        cands.append(_Candidate("layout", li, [s for s in shapes if s.ph_idx is not None], lay.name,
                                [s for s in shapes if s.ph_idx is None]))

    out: list[LayoutSpec] = []
    taken: set[str] = set()
    highlights = []
    design_shapes = []
    for pos, c in enumerate(cands):
        mapper = _Mapper(W, H, c.source == "layout")
        d = mapper.map(c.shapes, c.context)
        design_shapes.append(c.shapes)
        highlights += [s.highlight for s in d.texts if s.highlight]
        all_text = " ".join(s.text for s in [*c.shapes, *c.context] if s.kind == "text" and not s.decoration_text)
        n_pics = sum(1 for s in c.shapes if s.kind == "picture")
        slide_pos = pos if c.source == "slide" else -1
        slide_total = sum(1 for x in cands if x.source == "slide")
        kind, aliases, conf = _kind(d, slide_pos, slide_total, all_text, n_pics)
        hint = _name_hint(c.name, d)
        if hint and hint[0] != kind and conf < 0.8:
            kind, aliases, conf = hint
        if kind == "timeline" and items_has(d, "title") and not items_has(d, "label"):
            d.items = {("label" if k == "title" else k): v for k, v in d.items.items()}
            d.item_examples = [{("label" if k == "title" else k): v for k, v in e.items()} for e in d.item_examples]
        if kind in ("section", "title", "statement"):
            d.fields = [f.model_copy(update={"name": "subtitle"}) if f.name == "body" and
                        not any(x.name == "subtitle" for x in d.fields) else f for f in d.fields]
            if "body" in d.example and "subtitle" not in d.example:
                d.example["subtitle"] = d.example.pop("body")
        items = None
        if d.items:
            n = len(next(iter(d.items.values())))
            fixed = kind in ("cards", "image") or (c.source == "layout" and kind != "columns") or "card" in d.extra
            lo = n if fixed else (2 if kind in ("stats", "chart", "timeline") else 1)
            order = [k for k in ("number", "value", "image", "title", "label", "body", "caption") if k in d.items]
            order += [k for k in d.items if k not in order]
            content = [k for k in order if k not in d.auto]
            items = ItemsSpec(targets={k: d.items[k] for k in order}, min=min(lo, n), max=n,
                              required=(content[0],) if content else (), auto=d.auto, extra=d.extra,
                              types=d.item_types,
                              fixed_note=f"The layout shows {n} items. Provide exactly {n}." if lo == n else "")
            if d.item_examples:
                d.example["items"] = d.item_examples
        n_items = items.max if items else None
        lid = _slug(kind, n_items if kind not in ("title", "section", "statement", "text", "quote", "code",
                                                  "closing", "qa", "divider", "stat") else None, taken)
        desc, use = DESCRIPTIONS.get(kind, ("", ""))
        spec = LayoutSpec(
            id=lid, name=f"{kind.replace('-', ' ').capitalize()} ({c.source} {c.ref})", kind=kind,
            source=c.source, ref=c.ref, description=desc, use_when=use, fields=tuple(d.fields), items=items,
            handler=d.handler, aliases=aliases, example=_fill_examples(d.fields, items, d.example, {x.id: x for x in c.shapes}), confidence=conf,
            footer=bool(footer) and c.source == "layout",
        )
        out.append(spec)
        if conf < 0.6:
            issues.append(Issue(layout=lid, message=f"low confidence ({conf:.2f}) in kind '{kind}'"))
        if d.unmapped:
            issues.append(Issue(layout=lid, message=f"text shapes not mapped: {d.unmapped}"))

    brand = _brand(prs, ins, design_shapes, hidden_text, highlights)
    return Pack(
        id=pack_id,
        name=name or pptx.stem,
        description=f"Generated from {pptx.name}.",
        source_sha256=sha256(pptx),
        generator_version=GENERATOR_VERSION,
        status="draft",
        brand=brand,
        footer=footer,
        guide="",
        layouts=tuple(out),
        issues=tuple(issues),
    )
